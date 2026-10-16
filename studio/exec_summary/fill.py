"""Fill the Carrier Exec Summary template — a named step pipeline over one Presentation.

The template (``executive summary/Carrier Exec Summary Deck 1.pptx``) is fixed and
author-made: cover, summary, detail, back cover. Shapes are found by what they ARE (the
2-column table is the rank table, the wide one is the line-of-business table, charts in
left-to-right order) rather than by authored names, so a re-saved template still fills.
"""
from __future__ import annotations

import copy
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from studio.exec_summary import pptx_text as T
from studio.exec_summary.slides import ChartSeries, SlideContent
from studio.steps import PipelineBuilder

TEMPLATE_PATH = (Path(__file__).resolve().parents[2] / "executive summary"
                 / "Carrier Exec Summary Deck 1.pptx")

COVER, SUMMARY, DETAIL = 0, 1, 2
GREEN, RED = "14853D", "C53532"
NAVY = "000F47"
MIX_COLUMN_AT = 4            # Share of Portfolio goes before Share of Wallet
MIX_HEADER = "Share of Portfolio"


@dataclass
class FillContext:
    prs: object
    content: SlideContent
    carrier_label: str
    out_path: str


def build_fill_context(content: SlideContent, out_path: str, *, carrier_label: str,
                       template_path: Optional[str] = None) -> FillContext:
    return FillContext(prs=Presentation(str(template_path or TEMPLATE_PATH)), content=content,
                       carrier_label=carrier_label, out_path=out_path)


# ── shape lookup ─────────────────────────────────────────────────────────────


def _slide(ctx: FillContext, index: int):
    return ctx.prs.slides[index]


def _by_text(slide, starts: str):
    return next((s for s in slide.shapes if s.has_text_frame
                 and s.text_frame.text.strip().startswith(starts)), None)


def _title(slide):
    return next((s for s in slide.placeholders if s.placeholder_format.type == 1), None) \
        or _by_text(slide, "Carrier")


def _tables(slide) -> List:
    return [s for s in slide.shapes if s.has_table]


def _charts(slide) -> List:
    return sorted((s for s in slide.shapes if getattr(s, "has_chart", False) and s.has_chart),
                  key=lambda s: s.left)


#: The footer band: under the layout's rule (~6.40M EMU), above the MARSH logo (~6.54M).
#: A footnote there never collides with content, whatever height PowerPoint gives the
#: table once its header wraps.
FOOTNOTE_TOP = 6_420_000
FOOTNOTE_WIDTH = 9_400_000


def _add_footnote(slide, text: str, top: int = FOOTNOTE_TOP, width: int = FOOTNOTE_WIDTH) -> None:
    box = slide.shapes.add_textbox(Emu(457200), Emu(top), Emu(width), Emu(110000))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(7)
    T.set_color(frame.paragraphs[0], NAVY)


# ── steps ────────────────────────────────────────────────────────────────────


def write_cover(ctx: FillContext) -> None:
    slide = _slide(ctx, COVER)
    title = _by_text(slide, "Carrier Exec Summary")
    if title is not None:
        T.set_text(title.text_frame, ctx.content.cover_title)
    dated = _by_text(slide, "October XX")
    if dated is not None:
        T.set_text(dated.text_frame, ctx.content.cover_date)


def write_summary(ctx: FillContext) -> None:
    slide = _slide(ctx, SUMMARY)
    title = _title(slide)
    T.set_text(title.text_frame, ctx.content.summary_title)
    body = _by_text(slide, "Marsh-placed premium")
    text = ctx.content.summary_paragraph or "Commentary unavailable for this selection."
    T.set_text(body.text_frame, text)
    size = summary_font_size(text)
    if size is not None:
        for run in body.text_frame.paragraphs[0].runs:
            run.font.size = Pt(size)
    _add_footnote(slide, ctx.content.summary_footnote)


#: The summary panel holds about this many characters at the template's 12pt; past it,
#: a selection of many countries steps the type down rather than spilling off the panel.
SUMMARY_FIT_12PT = 3400
SUMMARY_FIT_11PT = 4000


def summary_font_size(text: str) -> Optional[int]:
    """``None`` to keep the template's 12pt, else the smaller size the paragraph needs."""
    if len(text) <= SUMMARY_FIT_12PT:
        return None
    return 11 if len(text) <= SUMMARY_FIT_11PT else 10


def write_detail_title(ctx: FillContext) -> None:
    T.set_text(_title(_slide(ctx, DETAIL)).text_frame, ctx.content.detail_title)


def _fill_chart(chart_shape, series: ChartSeries, carrier_label: str) -> None:
    data = CategoryChartData(number_format=series.number_format)
    data.categories = list(series.categories)
    data.add_series(carrier_label, list(series.carrier))
    data.add_series("Peer avg", list(series.peer))
    chart_shape.chart.replace_data(data)


def fill_charts(ctx: FillContext) -> None:
    slide = _slide(ctx, DETAIL)
    premium, sow = _charts(slide)[:2]
    _fill_chart(premium, ctx.content.premium_chart, ctx.carrier_label)
    _fill_chart(sow, ctx.content.sow_chart, ctx.carrier_label)
    for starts, unit in (("Carrier Premium vs Peers", ctx.content.premium_chart.unit),
                         ("Carrier SoW vs Peers", "%")):
        label = _by_text(slide, starts)
        if label is not None:
            T.set_text(label.text_frame, f"{starts} ({unit})")


def _trim_rows(table_shape, keep: int) -> None:
    """Remove body rows past ``keep`` and shrink the frame to the rows left."""
    tbl = table_shape.table._tbl
    for tr in tbl.tr_lst[keep + 1:]:
        tbl.remove(tr)
    table_shape.height = Emu(sum(int(tr.h) for tr in tbl.tr_lst))


def fill_rank_table(ctx: FillContext) -> None:
    table_shape = next(t for t in _tables(_slide(ctx, DETAIL)) if len(t.table.columns) == 2)
    rows = ctx.content.rank_rows
    for i, (country, rank) in enumerate(rows, start=1):
        if i >= len(table_shape.table.rows):
            break
        T.set_cell(table_shape.table.cell(i, 0), country)
        T.set_cell(table_shape.table.cell(i, 1), rank)
    _trim_rows(table_shape, len(rows))


def _insert_column(table_shape, at: int) -> None:
    """Clone column ``at`` (cells and grid) in front of itself, keeping the frame width."""
    tbl = table_shape.table._tbl
    grid = tbl.tblGrid
    cols = grid.findall(qn("a:gridCol"))
    total = sum(int(c.get("w")) for c in cols)
    cols[at].addprevious(copy.deepcopy(cols[at]))
    for tr in tbl.tr_lst:
        tr.tc_lst[at].addprevious(copy.deepcopy(tr.tc_lst[at]))
    cols = grid.findall(qn("a:gridCol"))
    first = int(cols[0].get("w")) * 9 // 10
    rest = (total - first) // (len(cols) - 1)
    cols[0].set("w", str(total - rest * (len(cols) - 1)))
    for col in cols[1:]:
        col.set("w", str(rest))


def _borrow_format(src_cell, dst_cell) -> None:
    """Give a blank authored cell the run format of a filled sibling in the same row."""
    if dst_cell.text_frame.paragraphs[0].runs:
        return
    body = dst_cell._tc.txBody
    for p in body.findall(qn("a:p")):
        body.remove(p)
    body.append(copy.deepcopy(src_cell.text_frame.paragraphs[0]._p))


def fill_lob_table(ctx: FillContext) -> None:
    table_shape = next(t for t in _tables(_slide(ctx, DETAIL)) if len(t.table.columns) >= 7)
    _insert_column(table_shape, MIX_COLUMN_AT)
    table = table_shape.table
    T.set_cell(table.cell(0, MIX_COLUMN_AT), MIX_HEADER)
    rows = ctx.content.table_rows
    last = len(table.columns) - 1
    for i, row in enumerate(rows, start=1):
        if i >= len(table.rows):
            break
        cells = [table.cell(i, c) for c in range(len(table.columns))]
        for blank in (1, last):
            _borrow_format(cells[2], cells[blank])
        values = (row.product, row.marsh, row.carrier, row.yoy, row.mix, row.sow, row.rank,
                  row.peers)
        for c, value in enumerate(values):
            T.set_cell(cells[c], value)
        T.set_color(cells[3].text_frame.paragraphs[0],
                    GREEN if row.yoy_sign > 0 else RED if row.yoy_sign < 0 else NAVY)
    _trim_rows(table_shape, len(rows))


def write_note(ctx: FillContext) -> None:
    slide = _slide(ctx, DETAIL)
    note = _by_text(slide, "Top3") or _by_text(slide, "Top 3")
    if note is not None:
        T.set_text(note.text_frame, ctx.content.note)
    _add_footnote(slide, ctx.content.detail_footnote)


def save_deck(ctx: FillContext) -> None:
    Path(ctx.out_path).parent.mkdir(parents=True, exist_ok=True)
    ctx.prs.save(ctx.out_path)


FILL_PIPELINE = (
    PipelineBuilder("fill_exec_summary")
    .step("write_cover", write_cover)
    .step("write_summary", write_summary)
    .step("write_detail_title", write_detail_title)
    .step("fill_charts", fill_charts)
    .step("fill_rank_table", fill_rank_table)
    .step("fill_lob_table", fill_lob_table)
    .step("write_note", write_note)
    .step("save_deck", save_deck)
    .build()
)


def fill_exec_summary(content: SlideContent, out_path: str, *, carrier_label: str,
                      template_path: Optional[str] = None) -> str:
    ctx = build_fill_context(content, out_path, carrier_label=carrier_label,
                             template_path=template_path)
    FILL_PIPELINE.run(ctx)
    return ctx.out_path
