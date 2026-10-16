"""Write text into authored shapes without losing the author's formatting.

The template's text is split across several runs (placeholder tokens like "xxxM" sit in
their own runs). Writing ``text_frame.text`` would throw every run's formatting away and
fall back to theme defaults, so these helpers keep the FIRST run's properties and replace
everything after it.
"""
from __future__ import annotations

import copy
from typing import Optional

from lxml import etree
from pptx.oxml.ns import qn

_KEEP_RPR_ATTRS = ("lang", "sz", "b", "i", "u", "dirty")


def _clean_rpr(rpr) -> None:
    """Drop proofing marks ("err") carried over from the authored placeholder run."""
    for attr in list(rpr.attrib):
        if attr not in _KEEP_RPR_ATTRS:
            del rpr.attrib[attr]


def set_paragraph_text(paragraph, text: str) -> None:
    """One paragraph's text, in the format of its first run (or field)."""
    p = paragraph._p
    first = p.find(qn("a:r"))
    if first is None:
        first = p.find(qn("a:fld"))
    rpr = copy.deepcopy(first.find(qn("a:rPr"))) if first is not None and first.find(qn("a:rPr")) is not None else None
    for child in list(p):
        if child.tag in (qn("a:r"), qn("a:fld"), qn("a:br")):
            p.remove(child)
    run = etree.SubElement(p, qn("a:r"))
    if rpr is not None:
        _clean_rpr(rpr)
        run.append(rpr)
    t = etree.SubElement(run, qn("a:t"))
    t.text = text
    end = p.find(qn("a:endParaRPr"))
    if end is not None:
        p.remove(end)
        p.append(end)


def set_text(text_frame, text: str) -> None:
    """The whole frame as ONE paragraph, in the first paragraph's format."""
    paragraphs = text_frame.paragraphs
    set_paragraph_text(paragraphs[0], text)
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def set_color(paragraph, hex_rgb: Optional[str]) -> None:
    """Recolour every run of a paragraph (a value cell's sign colour)."""
    if not hex_rgb:
        return
    for r in paragraph._p.findall(qn("a:r")):
        rpr = r.find(qn("a:rPr"))
        if rpr is None:
            rpr = etree.Element(qn("a:rPr"))
            r.insert(0, rpr)
        for fill in rpr.findall(qn("a:solidFill")):
            rpr.remove(fill)
        fill = etree.Element(qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", hex_rgb)
        # Schema order: an <a:ln> outline comes before the fill.
        rpr.insert(1 if rpr.find(qn("a:ln")) is not None else 0, fill)


def run_color(paragraph) -> Optional[str]:
    """The authored srgb colour of a paragraph's first run, if it has one."""
    r = paragraph._p.find(qn("a:r"))
    clr = r.find(f"{qn('a:rPr')}/{qn('a:solidFill')}/{qn('a:srgbClr')}") if r is not None else None
    return clr.get("val") if clr is not None else None


def set_cell(cell, text: str, color: Optional[str] = None) -> None:
    set_text(cell.text_frame, text)
    set_color(cell.text_frame.paragraphs[0], color)


def replace_in_runs(text_frame, old: str, new: str) -> bool:
    """Replace a token inside whichever run holds it (cover/title substitutions)."""
    hit = False
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
                hit = True
    return hit
