"""Typing straight on the Boardroom Canvas — commentary, titles and labels.

The canvas tags every authored line with a path into its widget's props; the
browser sends the retyped line back down the same sink the drag/resize actions
use. These tests walk that whole path: the rules, the rendered surface, the
action router, and the PowerPoint the edited document exports to.
"""
from __future__ import annotations

import pathlib

import pytest
from dash.development.base_component import Component
from pptx import Presentation

from studio.authoring.editing import CANVAS_ACTIONS
from studio.deck.model import ChartBlock, DeckSpec, SlideSpec
from studio.export import export_document
from studio.page import canvas as CV
from studio.page import document as D
from studio.page import inline_edit as IE


_TAKEAWAYS = (
    {"label": "Performance.", "text": "Premium grew ahead of plan.", "tone": "good"},
    {"label": "Watch item.", "text": "Property retention slipped.", "tone": "warn"},
    {"label": "Next.", "text": "Reprice the casualty book.", "tone": "neutral"},
)


def _deck() -> DeckSpec:
    """A generated deck with the commentary and chart a real QBR page carries."""
    return DeckSpec(
        slides=(
            SlideSpec(layout="cover", title="Zurich Singapore", eyebrow="QBR"),
            SlideSpec(
                layout="exec",
                title="Executive summary",
                eyebrow="SUMMARY",
                takeaways=_TAKEAWAYS,
            ),
            SlideSpec(
                layout="insight",
                title="Premium grew",
                eyebrow="PERFORMANCE",
                takeaways=_TAKEAWAYS,
                blocks=(
                    ChartBlock(
                        chart="bar",
                        labels=["Marine", "Property", "Casualty"],
                        values=[85.0, 75.0, 20.0],
                        title="Premium by product",
                    ),
                ),
            ),
        ),
        meta={"carrier": "Zurich", "country": "Singapore", "year": 2025},
    )


def _walk(node):
    if node is None:
        return
    if isinstance(node, (list, tuple)):
        for child in node:
            yield from _walk(child)
        return
    yield node
    if isinstance(node, Component):
        yield from _walk(getattr(node, "children", None))


def _editables(node):
    """Every text node the browser is allowed to retype, in render order."""
    return [
        item for item in _walk(node)
        if getattr(item, "contentEditable", None) == "true"
    ]


def _editable_paths(node) -> list[str]:
    return [getattr(item, "data-qs-path") for item in _editables(node)]


def _commentary(doc, sid):
    return next(w for w in D.page_widgets(doc, sid) if w["kind"] == "text")


# ── the rules: what carries a path at all ────────────────────────────────────


@pytest.mark.parametrize(
    "path,expected",
    [
        ("title", ("title",)),
        ("text", ("text",)),
        ("points.0.text", ("points", 0, "text")),
        ("items.11.owner", ("items", 11, "owner")),
    ],
)
def test_an_authored_path_resolves_to_one_leaf(path, expected):
    assert IE.parse_path(path) == expected


@pytest.mark.parametrize(
    "path",
    [
        "values",            # a data series, not prose
        "font_color",        # appearance, owned by the inspector
        "points",            # a container, not a leaf
        "points.text",       # missing index
        "points.x.text",     # index is not a number
        "rows.0.premium",    # table data
        "props.points.0.text",
        "",
        None,
    ],
)
def test_a_data_or_malformed_path_carries_no_edit(path):
    assert IE.parse_path(path) is None


def test_emptying_a_commentary_bullet_removes_it():
    props = {"points": [{"text": "Kept"}, {"text": "Dropped"}, {"text": "Also kept"}]}

    out = IE.apply_text(props, IE.parse_path("points.1.text"), "   ")

    assert [p["text"] for p in out["points"]] == ["Kept", "Also kept"]
    assert props["points"][1]["text"] == "Dropped", "the input must not be mutated"


def test_an_out_of_range_bullet_leaves_the_commentary_alone():
    props = {"points": [{"text": "Only one"}]}

    assert IE.apply_text(props, IE.parse_path("points.4.text"), "Ghost") == props


def test_a_new_bullet_opens_directly_below_the_one_being_written():
    props = {"points": [{"text": "First"}, {"text": "Second"}]}

    out = IE.add_point(props, 0)

    assert [p["text"] for p in out["points"]] == ["First", "", "Second"]
    assert out["points"][1]["tone"] == "neutral"


# ── committing an edit to the shared document ────────────────────────────────


def test_a_retyped_commentary_point_lands_in_the_document():
    doc = D.new_document(_deck())
    sid = D.sid_at(doc, 1)
    commentary = _commentary(doc, sid)

    doc = D.set_widget_text(
        doc, sid, commentary["id"], "points.0.text", "Marine grew 18% on the year"
    )

    points = D.get_widget(doc, sid, commentary["id"])["props"]["points"]
    assert points[0]["text"] == "Marine grew 18% on the year"
    assert points[0]["label"], "retyping the sentence keeps the bullet's label"


def test_a_path_the_rules_reject_cannot_rewrite_the_document():
    doc = D.new_document(_deck())
    sid = D.sid_at(doc, 1)
    commentary = _commentary(doc, sid)

    for path in ("tone", "points.0.tone", "font_color", "points"):
        assert D.set_widget_text(doc, sid, commentary["id"], path, "hacked") == doc


def test_a_headline_typed_on_the_canvas_becomes_the_page_title():
    doc = D.new_document(_deck())
    sid = D.sid_at(doc, 1)
    headline = next(w for w in D.page_widgets(doc, sid) if w["kind"] == "headline")

    doc = D.set_widget_text(doc, sid, headline["id"], "text", "Growth beat plan")
    doc = D.set_widget_text(doc, sid, headline["id"], "subtitle", "Q2 2025")

    props = D.get_widget(doc, sid, headline["id"])["props"]
    assert props["text"] == "Growth beat plan"
    assert props["subtitle"] == "Q2 2025"


# ── the rendered surface ─────────────────────────────────────────────────────


def test_the_selected_widget_offers_every_authored_line_for_retyping():
    doc = D.new_document(_deck())
    sid = D.sid_at(doc, 1)
    commentary = _commentary(doc, sid)
    widgets = D.page_widgets(doc, sid)

    paths = _editable_paths(CV.canvas_surface(widgets, commentary["id"]))
    points = commentary["props"]["points"]

    assert "heading" in paths
    for index in range(len(points)):
        assert f"points.{index}.text" in paths
        assert f"points.{index}.label" in paths


def test_an_unselected_widget_and_a_thumbnail_stay_read_only():
    doc = D.new_document(_deck())
    sid = D.sid_at(doc, 1)
    widgets = D.page_widgets(doc, sid)
    commentary = _commentary(doc, sid)

    surface = CV.canvas_surface(widgets, commentary["id"])
    owners = {getattr(item, "data-qs-wid") for item in _editables(surface)}

    assert owners == {commentary["id"]}, "only the selected widget is editable"
    assert _editable_paths(CV.canvas_surface(widgets, None)) == []
    assert _editable_paths(CV.thumbnail_surface(widgets)) == []


def test_a_chart_offers_its_title_even_when_the_generator_left_it_blank():
    doc = D.new_document(_deck())
    sid = D.sid_at(doc, 2)
    widgets = D.page_widgets(doc, sid)
    chart = next((w for w in widgets if w["kind"] == "chart"), None)
    assert chart is not None, "the insight page is the one with a chart"
    doc = D.set_widget_text(doc, sid, chart["id"], "title", "")

    surface = CV.canvas_surface(D.page_widgets(doc, sid), chart["id"])

    assert "title" in _editable_paths(surface)


def test_editable_text_is_remounted_when_the_stored_text_changes():
    """React owns the node, so a changed key is what flushes a reverted edit."""
    doc = D.new_document(_deck())
    sid = D.sid_at(doc, 1)
    headline = next(w for w in D.page_widgets(doc, sid) if w["kind"] == "headline")

    def title_key(document):
        surface = CV.canvas_surface(D.page_widgets(document, sid), headline["id"])
        return next(
            getattr(item, "key")
            for item in _walk(surface)
            if getattr(item, "data-qs-path", None) == "text"
        )

    before = title_key(doc)
    after = title_key(D.set_widget_text(doc, sid, headline["id"], "text", "Rewritten"))

    assert before != after


# ── the action router the browser talks to ───────────────────────────────────


def test_the_canvas_routes_a_typed_line_and_keeps_the_widget_selected():
    doc = D.new_document(_deck())
    sid = D.sid_at(doc, 1)
    commentary = _commentary(doc, sid)

    updated, selected = CANVAS_ACTIONS["text"](
        doc,
        sid,
        {"wid": commentary["id"], "path": "points.0.text", "value": "Retyped", "then": ""},
    )

    assert selected == commentary["id"]
    assert D.get_widget(updated, sid, commentary["id"])["props"]["points"][0]["text"] == "Retyped"


def test_pressing_enter_commits_the_line_and_opens_the_next_bullet():
    doc = D.new_document(_deck())
    sid = D.sid_at(doc, 1)
    commentary = _commentary(doc, sid)
    before = len(commentary["props"]["points"])

    updated, _ = CANVAS_ACTIONS["text"](
        doc,
        sid,
        {
            "wid": commentary["id"],
            "path": "points.0.text",
            "value": "Marine led the growth",
            "then": "point-add",
        },
    )

    points = D.get_widget(updated, sid, commentary["id"])["props"]["points"]
    assert len(points) == before + 1
    assert points[0]["text"] == "Marine led the growth"
    assert points[1]["text"] == ""


def test_an_unknown_canvas_action_is_simply_not_routed():
    assert CANVAS_ACTIONS.get("drop-database") is None
    assert set(CANVAS_ACTIONS) == {"select", "geo", "text"}


# ── the wiring that makes it work in the browser ─────────────────────────────


def test_the_canvas_script_reads_the_path_and_leaves_the_caret_alone():
    js = (pathlib.Path(__file__).parents[1] / "assets/studio_canvas.js").read_text(
        encoding="utf-8"
    )

    assert 'var EDITABLE = "[data-qs-path]";' in js
    # a click on text must not preventDefault, or the caret never lands
    assert "if (closest(e.target, EDITABLE)) return;" in js
    assert 'action: "text"' in js
    assert '"point-add"' in js


def test_the_editable_affordance_is_styled():
    css = (pathlib.Path(__file__).parents[1] / "assets/studio_authoring.css").read_text(
        encoding="utf-8"
    )

    assert ".qs-cv-editable" in css
    assert "content: attr(data-qs-ph)" in css


# ── end to end: setup → canvas edit → materialize → PowerPoint ───────────────


def test_every_line_retyped_on_the_canvas_reaches_the_exported_deck(tmp_path):
    deck = _deck()
    doc = D.new_document(deck)
    sid = D.sid_at(doc, 1)
    headline = next(w for w in D.page_widgets(doc, sid) if w["kind"] == "headline")
    commentary = _commentary(doc, sid)

    # what a user does on the slide: retype the title, rewrite bullet one,
    # press Enter, write the new bullet, then blank bullet three away.
    doc = D.set_widget_text(doc, sid, headline["id"], "text", "Marine carried the quarter")
    doc, _ = CANVAS_ACTIONS["text"](
        doc,
        sid,
        {
            "wid": commentary["id"],
            "path": "points.0.text",
            "value": "Marine premium grew 18% year on year",
            "then": "point-add",
        },
    )
    doc = D.set_widget_text(
        doc, sid, commentary["id"], "points.1.text", "Property retention needs a plan"
    )
    doc = D.set_widget_text(doc, sid, commentary["id"], "points.2.text", "")

    materialized = D.materialize(doc, for_export=True)
    assert materialized is not None

    out = tmp_path / "canvas-edited.pptx"
    export_document(doc, out_path=str(out))
    slide = Presentation(out).slides[1]
    text = "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))

    assert "Marine carried the quarter" in text
    assert "Marine premium grew 18% year on year" in text
    assert "Property retention needs a plan" in text
    assert "Executive summary" not in text, "the retyped title replaced the generated one"
