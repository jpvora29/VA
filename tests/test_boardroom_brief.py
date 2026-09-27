"""The Boardroom as an executive brief, and a board with no invented ratings.

Two things are pinned here.

**The layout.** The chapters are numbered tabs above the pages (the old slider sat
below them, where a reader only found it after scrolling past the first chapter);
every chapter states its question as the headline and carries its scope line; the
first KPI is the headline figure; and each chapter's footer names the way back and
the way on.

**Nothing interpreted.** Every figure on the board is a measure a reader can
check. The model's ratings — 0-100 "gap scores", heatmap intensities, a 2x2 on
made-up axes, High/Med/Low risk severities, "Low presence / Established" pills and
its good/bad colouring — are never drawn, on screen or on the slide.

The business flow is covered end to end: digest -> document -> rendered board ->
exported PPTX, for a saved board that still holds the retired widgets.

Run:  pytest tests/test_boardroom_brief.py -q -o pythonpath=.
"""
from __future__ import annotations

import io

from dash import dcc
from dash.development.base_component import Component
from pptx import Presentation

from ui.boardroom import builder, catalog, editor, ppt_export, retired, widgets_generated
from ui.boardroom.render import render_document
from ui.boardroom.widgets_library import delta_direction


def walk(node):
    if isinstance(node, Component):
        yield node
        for child in node._traverse():
            if isinstance(child, Component):
                yield child
    elif isinstance(node, (list, tuple)):
        for item in node:
            yield from walk(item)


def text_of(node) -> str:
    """Visible text only: children strings, not tooltips."""
    parts = []
    for component in walk(node):
        children = getattr(component, "children", None)
        if isinstance(children, str):
            parts.append(children)
    return " ".join(parts)


def with_class(node, cls: str):
    return [n for n in walk(node) if cls in (getattr(n, "className", "") or "").split()]


def deck_text(doc) -> str:
    prs = Presentation(io.BytesIO(ppt_export.export_pptx(doc, [])))
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    text.extend(cell.text for cell in row.cells)
    return "\n".join(text)


def digest() -> dict:
    return {
        "title": "Zurich — Canada",
        "subtitle": "Q2 2026 premium performance",
        "headline": "Property contraction cost $3.2M in the quarter.",
        "kpis": [
            {"label": "Premium", "value": "$19.2M", "delta": "-14.3% QoQ", "tone": "good"},
            {"label": "Share of wallet", "value": "8.2%", "delta": "+0.4 pp", "tone": "danger"},
        ],
        "insights": [
            {"figure": "62%", "headline": "Property carries the book",
             "detail": "Of total premium.", "tone": "danger"},
        ],
        "commentary": [{"heading": "The big picture", "points": ["Premium fell in Property."]}],
        "risks": [
            {"label": "Rank decline", "severity": "High", "tone": "danger"},
            {"label": "Property contraction", "evidence": "Premium -$3.2M (-14.3% QoQ)"},
        ],
        "scope": [
            {"key": "country", "label": "Country", "value": "Canada", "source": "asked"},
            {"key": "product", "label": "Product", "value": "Property", "source": "asked"},
        ],
        "top_carriers": {
            "carriers": [
                {"carrier": "Zurich", "is_subject": True, "rank": 3, "premium": "$19.2M",
                 "premium_value": 19_200_000, "share_of_wallet_pct": 8.2},
                {"carrier": "Peer 1", "rank": 1, "premium": "$40.0M", "premium_value": 40_000_000},
            ],
        },
        "headroom": {
            "rows": [
                {"product_line": "Property", "carrier_premium_value": 8_200_000,
                 "marsh_premium_value": 42_000_000, "share_of_wallet_pct": 19.5,
                 "status": "established"},
                {"product_line": "Cyber", "carrier_premium_value": 0.0,
                 "marsh_premium_value": 9_000_000, "share_of_wallet_pct": 0.0,
                 "status": "low_presence"},
            ],
        },
    }


def saved_board_with_retired_widgets() -> dict:
    """A board saved before the score widgets were retired."""
    payload = digest()
    payload.update(
        {
            "opportunities": [
                {"area": "Cyber — Canada", "dimension": "product", "carrier_level": "$0.4M",
                 "peer_level": "$12.0M", "gap_score": 87, "recommendation": "Enter via MGA",
                 "tone": "good"},
            ],
            "opportunity_map": {
                "rows": ["Cyber"], "cols": ["Canada"],
                "cells": [{"row": "Cyber", "col": "Canada", "intensity": 73, "tone": "good",
                           "note": "Marsh book growing"}],
                "legend": "Darker = higher priority",
            },
        }
    )
    return payload


# ── the layout ───────────────────────────────────────────────────────────────


def test_chapters_are_tabs_above_the_pages_not_a_slider_below_them():
    doc = builder.build_document_from_digest(digest(), 0)
    rendered = render_document(doc, [], card_idx=4)
    assert not [n for n in walk(rendered) if isinstance(n, dcc.Slider)]

    tabs = with_class(rendered, "bm-chapter")
    assert [t.id["page"] for t in tabs] == list(range(len(doc["pages"])))
    assert all(t.id["idx"] == 4 for t in tabs)
    assert ["is-active" in t.className for t in tabs] == [True] + [False] * (len(tabs) - 1)

    # The tab bar comes before the pages in the card.
    order = [c.className for c in rendered.children]
    assert order.index("bm-chapter-bar") < order.index("bm-pages")


def test_an_edit_re_render_reopens_the_chapter_the_reader_was_on():
    doc = builder.build_document_from_digest(digest(), 0)
    rendered = render_document(doc, [], active_page=1)
    tabs = with_class(rendered, "bm-chapter")
    assert "is-active" in tabs[1].className
    pages = [n for n in walk(rendered) if isinstance(getattr(n, "id", None), dict)
             and n.id.get("type") == "bm-page"]
    assert [p.style for p in pages][:2] == [{"display": "none"}, {}]


def test_every_chapter_states_its_question_and_its_scope():
    doc = builder.build_document_from_digest(digest(), 0)
    rendered = render_document(doc, [])
    headlines = [h.children for h in with_class(rendered, "bm-brief-headline")]
    assert headlines == [p["caption"] for p in doc["pages"]]
    scope_lines = with_class(rendered, "bm-scope-line")
    assert len(scope_lines) == len(doc["pages"])
    line = text_of(scope_lines[0])
    assert "Zurich — Canada" in line and "Canada" in line and "Property" in line


def test_each_chapter_footer_names_the_way_back_and_the_way_on():
    doc = builder.build_document_from_digest(digest(), 0)
    titles = [p["title"] for p in doc["pages"]]
    rendered = render_document(doc, [])
    feet = with_class(rendered, "bm-brief-foot")
    assert len(feet) == len(titles)

    first, last = feet[0], feet[-1]
    assert f"1 of {len(titles)} — {titles[0]}" in text_of(first)
    assert f"Next: {titles[1]}" in text_of(first)
    assert not with_class(first, "prev")
    assert not with_class(last, "next") and with_class(last, "prev")
    steps = with_class(first, "bm-foot-step")
    assert steps[0].id == {"type": "bm-goto", "idx": 0, "page": 1, "src": 0}


def test_the_first_kpi_is_the_headline_figure():
    rendered = render_document(builder.build_document_from_digest(digest(), 0), [])
    hero = with_class(rendered, "is-hero")
    assert len(hero) == 1 and "$19.2M" in text_of(hero[0])
    assert len(with_class(rendered, "bm-kpi-rest")) == 1


def test_a_kpi_change_is_coloured_by_its_sign_not_the_models_opinion():
    """The fixture's tone says a 14% fall is 'good'; the board says what the sign says."""
    rendered = render_document(builder.build_document_from_digest(digest(), 0), [])
    deltas = with_class(rendered, "bm-kpi-delta")
    assert "down" in deltas[0].className.split()
    assert "up" in deltas[1].className.split()
    assert delta_direction("+12% YoY") == "up"
    assert delta_direction("−1.2 pp") == "down"
    assert delta_direction("vs Q1 2026") == ""


# ── nothing interpreted ──────────────────────────────────────────────────────


def test_a_watch_item_states_its_figure_and_never_a_severity():
    payload = {k: v for k, v in digest().items() if k != "watchlist"}
    rendered = render_document(builder.build_document_from_digest(payload, 0), [])
    shown = text_of(rendered)
    assert "Premium -$3.2M (-14.3% QoQ)" in shown
    assert "High" not in shown.split()
    assert not with_class(rendered, "bm-risk-track"), "the severity bar is gone"


def test_an_insight_is_led_by_the_figure_that_proves_it():
    rendered = render_document(builder.build_document_from_digest(digest(), 0), [])
    figures = with_class(rendered, "bm-insight-figure")
    assert [f.children for f in figures] == ["62%"]
    # No sentiment colour class on the card.
    card = with_class(rendered, "bm-insight-card")[0]
    assert not {"good", "warn", "danger"} & set(card.className.split())


def test_presence_is_read_from_the_premium_not_the_models_label():
    rendered = render_document(builder.build_document_from_digest(digest(), 0), [])
    pills = [p.children for p in with_class(rendered, "bm-x-status")]
    assert pills.count("No current premium") == 1, "only Cyber writes zero premium"
    assert "Low presence" not in pills and "Established" not in pills


def test_a_saved_score_widget_shows_its_figures_and_never_its_score():
    doc = builder.build_document_from_digest(saved_board_with_retired_widgets(), 0)
    kinds = {w["kind"] for p in doc["pages"] for w in p["widgets"]}
    assert {"opportunity_radar", "opportunity_map"} <= kinds

    shown = text_of(render_document(doc, []))
    assert "$0.4M" in shown and "$12.0M" in shown and "Enter via MGA" in shown
    assert "Marsh book growing" in shown
    for score in ("87", "73"):
        assert score not in shown.split(), f"the retired score {score} is still drawn"
    assert "Darker = higher priority" not in shown

    deck = deck_text(doc)
    assert "$12.0M" in deck and "Marsh book growing" in deck
    for score in ("87", "73"):
        assert score not in deck.split(), f"the slide still prints the score {score}"


def test_the_retired_rows_are_pure_and_keep_the_saved_order():
    rows = retired.evidence_rows(
        "opportunity_radar",
        {"opportunities": [{"area": "B", "gap_score": 10}, {"area": "A", "gap_score": 90}]},
    )
    assert [r.title for r in rows] == ["B", "A"], "sorting by the score would reinstate it"
    assert retired.evidence_rows("headroom", {"rows": []}) == []


def test_the_library_no_longer_offers_a_score_widget():
    offered = {w["kind"] for items in catalog.library_by_category().values() for w in items}
    assert not (offered & catalog.NOT_OFFERED)
    # ...but a saved one still renders and still opens in the editor.
    for kind in catalog.NOT_OFFERED:
        data = catalog.LIBRARY_BY_KIND[kind]["default_data"]
        assert isinstance(widgets_generated.render_bespoke(kind, data), Component)


def test_no_editor_offers_a_rating_field():
    banned = {"severity", "gap_score", "intensity", "premium_strength", "broker_perception",
              "status", "tone", "tones"}
    for spec in editor.LIST_SPECS.values():
        if spec["path"] == "battlecards":
            continue  # retired; its `broker_perception` is saved free text, not a score
        keys = {f["key"] for f in spec["fields"]}
        assert not (keys & banned), f"{spec['path']} still edits a rating: {keys & banned}"


def test_the_slide_prints_the_same_figures_and_no_ratings():
    payload = {k: v for k, v in digest().items() if k != "watchlist"}
    deck = deck_text(builder.build_document_from_digest(payload, 0))
    assert "Premium -$3.2M (-14.3% QoQ)" in deck
    assert "62%" in deck
    for rating in ("High", "Low presence", "Established"):
        assert rating not in deck, f"the slide still prints {rating!r}"
