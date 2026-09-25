"""Recap's own reasoning effort, token budgets and usage summary.

    RECAP_REASONING_EFFORT[_<STAGE>]  ->  settings.reasoning_effort_for(stage)
                                      ->  LLMClient.call(reasoning_effort=...)
                                      ->  _bind: the tier config the request is built from

The effort must never collide with the application-wide ``<TIER>_EFFORT`` knobs: an
unset Recap variable leaves every call exactly as its tier is configured, and the
chatbot's variables never leak into Recap's per-stage resolution.
"""
from __future__ import annotations

import asyncio
import logging
from types import SimpleNamespace

import pytest

from core.llm.clients import TierConfig
from recap.config import settings
from recap.generation.recap_generator import RecapGenerator
from recap.llm import llm_client
from recap.llm.llm_client import LLMClient, reasoning_tokens
from recap.metadata import read_cover
from recap.schemas.recap import TitledTakeaway

RECAP_VARS = ("RECAP_REASONING_EFFORT", "RECAP_REASONING_EFFORT_RECAP_DEDUP",
              "RECAP_TOKENS_ENRICHMENT_KPI", "REASONING_EFFORT", "BALANCED_EFFORT",
              "REASONING_EFFORT_RECAP_DEDUP", "TOKENS_ENRICHMENT_KPI")


@pytest.fixture(autouse=True)
def clean_env(monkeypatch):
    for name in RECAP_VARS:
        monkeypatch.delenv(name, raising=False)


# ── settings: which effort and budget a stage gets ──────────────────────────


def test_no_recap_effort_leaves_the_stage_to_its_tier():
    assert settings.reasoning_effort_for("recap_dedup") is None


def test_the_recap_wide_effort_applies_to_every_stage(monkeypatch):
    monkeypatch.setenv("RECAP_REASONING_EFFORT", "Medium")
    assert settings.reasoning_effort_for("recap_dedup") == "medium"
    assert settings.reasoning_effort_for("enrichment_kpi") == "medium"


def test_a_stage_effort_wins_over_the_recap_wide_one(monkeypatch):
    monkeypatch.setenv("RECAP_REASONING_EFFORT", "medium")
    monkeypatch.setenv("RECAP_REASONING_EFFORT_RECAP_DEDUP", "low")
    assert settings.reasoning_effort_for("recap_dedup") == "low"
    assert settings.reasoning_effort_for("recap_takeaway") == "medium"


@pytest.mark.parametrize("value", ["", "none", "off", "  "])
def test_an_empty_or_off_effort_means_unset(monkeypatch, value):
    monkeypatch.setenv("RECAP_REASONING_EFFORT", value)
    assert settings.reasoning_effort_for("recap_dedup") is None


def test_the_unprefixed_standalone_variables_are_not_read(monkeypatch):
    """The standalone app read REASONING_EFFORT / TOKENS_<STAGE>. In the merged app
    those names are too generic to own, so Recap reads only its RECAP_ ones."""
    monkeypatch.setenv("REASONING_EFFORT", "high")
    monkeypatch.setenv("REASONING_EFFORT_RECAP_DEDUP", "high")
    monkeypatch.setenv("TOKENS_ENRICHMENT_KPI", "9")
    assert settings.reasoning_effort_for("recap_dedup") is None
    assert settings.token_budget("enrichment_kpi") == 1024


def test_the_tier_effort_does_not_leak_into_the_recap_effort(monkeypatch):
    monkeypatch.setenv("BALANCED_EFFORT", "high")
    assert settings.reasoning_effort_for("umbrella_classification") is None


def test_token_budgets_have_defaults_and_overrides(monkeypatch):
    assert settings.token_budget("enrichment_kpi") == 1024
    assert settings.token_budget("no_such_stage") == 1024
    monkeypatch.setenv("RECAP_TOKENS_ENRICHMENT_KPI", "1500")
    assert settings.token_budget("enrichment_kpi") == 1500
    monkeypatch.setenv("RECAP_TOKENS_ENRICHMENT_KPI", "lots")
    assert settings.token_budget("enrichment_kpi") == 1024


# ── _bind: the request a call is built from ─────────────────────────────────


class FakeClient:
    """Records the config it was built from and every bind() on top of it."""

    def __init__(self, config: TierConfig) -> None:
        self.config = config
        self.bound: dict = {}

    def bind(self, **kwargs):
        self.bound.update(kwargs)
        return self


@pytest.fixture()
def bind_with(monkeypatch):
    """_bind against a fake factory, with the tier resolving to ``tier_config``."""

    def run(tier_config: TierConfig, budget: int, effort=None) -> FakeClient:
        import core.llm.clients as clients

        monkeypatch.setattr(clients, "resolve_tier", lambda tier: tier_config)
        monkeypatch.setattr(clients, "client_for", FakeClient)
        return llm_client._bind("balanced", budget, effort)

    return run


CLASSIC = TierConfig("gpt-classic", temperature=0.0)
REASONING = TierConfig("gpt-reasoning", effort="high")


def test_a_classic_tier_without_recap_effort_is_unchanged(bind_with):
    client = bind_with(CLASSIC, 512)
    assert client.config == CLASSIC
    assert client.bound == {"response_format": {"type": "json_object"}, "max_tokens": 512}


def test_a_recap_effort_turns_the_call_into_a_reasoning_request(bind_with):
    """Effort replaces temperature — a reasoning model rejects one — and the cap goes
    out as max_completion_tokens, the parameter a reasoning deployment accepts."""
    client = bind_with(CLASSIC, 512, "low")
    assert client.config.effort == "low"
    assert client.config.temperature is None
    assert client.config.deployment == "gpt-classic"
    assert client.bound["max_completion_tokens"] == 512
    assert "max_tokens" not in client.bound


def test_a_recap_effort_replaces_the_tier_effort(bind_with):
    client = bind_with(REASONING, 900, "minimal")
    assert client.config.effort == "minimal"


def test_a_reasoning_tier_without_recap_effort_keeps_its_own_effort_and_no_cap(bind_with):
    """The tier's effort is the app's setting; Recap does not cap its reasoning."""
    client = bind_with(REASONING, 900)
    assert client.config == REASONING
    assert "max_tokens" not in client.bound
    assert "max_completion_tokens" not in client.bound


# ── the usage summary ───────────────────────────────────────────────────────


class FakeRunnable:
    """Stands in for a bound LangChain client: answers with fixed usage."""

    def __init__(self, reasoning: int) -> None:
        self._reasoning = reasoning

    async def ainvoke(self, messages):
        return SimpleNamespace(content='{"ok": true}', usage_metadata={
            "input_tokens": 100, "output_tokens": 40, "total_tokens": 140,
            "output_token_details": {"reasoning": self._reasoning},
        })


def test_usage_is_summed_per_stage_and_cleared_by_the_summary(monkeypatch, caplog):
    monkeypatch.setattr(llm_client, "_bind",
                        lambda tier, budget, effort=None: FakeRunnable(30 if effort else 0))
    client = LLMClient()

    async def calls():
        await client.call(system_prompt="s", user_message="u", stage="recap_dedup",
                          reasoning_effort="low")
        await client.call(system_prompt="s", user_message="u", stage="recap_dedup",
                          reasoning_effort="low")
        await client.call(system_prompt="s", user_message="u")

    asyncio.run(calls())
    totals = client.get_usage_totals()
    assert totals["recap_dedup"] == {"calls": 2, "prompt_tokens": 200,
                                     "completion_tokens": 80, "total_tokens": 280,
                                     "reasoning_tokens": 60}
    assert totals["untagged"]["calls"] == 1

    with caplog.at_level(logging.INFO, logger=llm_client.__name__):
        client.log_usage_summary()
    assert "Takeaway Dedup ..... 60" in caplog.text
    assert "ALL stages ..... 60 (calls=3" in caplog.text
    assert client.get_usage_totals() == {}, "a summary covers calls since the last one"


def test_two_clients_do_not_share_usage(monkeypatch):
    """Two runs in one process — two users of the merged app — keep separate totals."""
    monkeypatch.setattr(llm_client, "_bind", lambda *a, **k: FakeRunnable(5))
    first, second = LLMClient(), LLMClient()
    asyncio.run(first.call(system_prompt="s", user_message="u", stage="recap_takeaway"))
    assert second.get_usage_totals() == {}


def test_reasoning_tokens_reads_langchain_usage():
    assert reasoning_tokens({"output_token_details": {"reasoning": 12}}) == 12
    assert reasoning_tokens({"output_token_details": {}}) == 0
    assert reasoning_tokens({}) == 0


# ── merged bullets keep their taxonomy without the model naming it ──────────


def test_merged_fields_come_from_the_source_bullets():
    takeaways = [
        TitledTakeaway(title="A", narrative="a", umbrella="performance_and_position",
                       sub_category="overall_trading_performance",
                       source_content_unit_ids=["cu1"]),
        TitledTakeaway(title="B", narrative="b", umbrella="performance_and_position",
                       sub_category="overall_trading_performance",
                       source_content_unit_ids=["cu2"]),
        TitledTakeaway(title="C", narrative="c", umbrella="opportunity_and_growth",
                       source_content_unit_ids=["cu3"]),
    ]
    generator = RecapGenerator(llm_client=object())
    umbrella, sub_category, ids = generator._derive_merged_fields([1, 2, 3], takeaways)
    assert umbrella == "performance_and_position"
    assert sub_category == "overall_trading_performance"
    assert ids == ["cu1", "cu2", "cu3"]


# ── the cover slide's scope chips ───────────────────────────────────────────


def test_the_cover_names_the_carrier_and_the_country(tmp_path):
    from pptx import Presentation

    carrier = settings.carriers[0]
    country = settings.countries[0]
    presentation = Presentation()
    cover = presentation.slides.add_slide(presentation.slide_layouts[0])
    cover.shapes.title.text = "Acme Corp QBR Q2 2026"
    cover.placeholders[1].text = f"{carrier} | {country}"
    path = tmp_path / "cover.pptx"
    presentation.save(path)

    metadata = read_cover(path.read_bytes())
    assert metadata.carrier == carrier
    assert metadata.country_region == country
