"""Every page follows the selection — its years, currency, period names and charts.

The regressions: the portfolio pages said "Marsh GWP 2025" / "2025 Rank" whatever year
was selected (their years sit in TABLE cells, which the template-year detector never
read); the cover kept the template's example date; "Opportunities for …" shifted off by
one; the GWP page printed euros on a dollar book, named its periods CY / PY, and drew its
two total bars from the template's example book beside the real totals; a one-country
portfolio page kept an empty second country row; a Casualty page carried Marine's ship.
"""
from __future__ import annotations

import pathlib
import re
from datetime import date

import pytest

from studio.template_fill import feedback, product_icon
from studio.template_fill.render import currency, in_reporting_currency
from studio.template_fill.year_labels import forward_year, template_year, year_subs

TEMPLATES = pathlib.Path(__file__).parents[1] / "template"
TODAY = date(2026, 9, 28)


def _relabel(text: str, **values) -> str:
    for pattern, repl in year_subs(values, today=TODAY):
        text = pattern.sub(repl, text)
    return text


def _template(name: str):
    from studio.template_fill.analyze import analyze

    path = TEMPLATES / name
    if not path.exists():
        pytest.skip(f"{name} is not checked out")
    return analyze(str(path))


# ── which year a template was written for ────────────────────────────────────


def test_years_in_table_cells_count_toward_the_template_year():
    """The product template's years ALL sit in its KPI table — none in a text box."""
    from studio.template_fill.model import _template_year

    assert _template_year(_template("product_template.pptx")) == 2025


def test_dates_copyright_and_forward_years_are_not_the_reporting_year():
    texts = ["FY 2025", "Opportunities for 2026", "22 July 2026", "© 2026 Marsh"]

    assert template_year(texts) == 2025


def test_a_tie_reports_on_the_later_period():
    assert template_year(["TTM April 2026", "TTM April 2025"]) == 2026


# ── how each kind of year is rewritten ───────────────────────────────────────

CALENDAR_2024 = dict(template_year=2025, period_year=2024, forward_year=2025)


@pytest.mark.parametrize("authored,expected", [
    ("Marsh GWP 2025", "Marsh GWP 2024"),
    ("2025 Rank", "2024 Rank"),
    ("Carrier trading highlights FY 2025", "Carrier trading highlights FY 2024"),
    ("TTM April 2025", "TTM April 2024"),
    ("Opportunities for 2026", "Opportunities for 2025"),
    ("22 July 2026", "28 September 2026"),
    ("© 2026 Marsh. All rights reserved.", "© 2026 Marsh. All rights reserved."),
])
def test_each_kind_of_year_follows_the_selection(authored, expected):
    assert _relabel(authored, **CALENDAR_2024) == expected


def test_a_rewritten_year_is_never_shifted_twice():
    """One pass: the build date and the forward year must not then move as reporting years."""
    assert _relabel("22 July 2026 · for 2026 · FY 2025", **CALENDAR_2024) == \
        "28 September 2026 · for 2025 · FY 2024"


def test_a_rolling_deck_names_periods_and_plans_twelve_months_ahead():
    values = dict(template_year=2025, period_year=2026,
                  period_labels={"2026": "TTM Aug 2026", "2025": "TTM Aug 2025"})

    assert _relabel("FY 2025", **values) == "TTM Aug 2026"
    assert _relabel("Marsh GWP 2024", **values) == "Marsh GWP TTM Aug 2025"
    assert _relabel("Opportunities for 2026", **values) == \
        "Opportunities for the next twelve months"


def test_the_basis_heading_follows_a_rolling_deck():
    from studio.template_fill.year_labels import basis_subs

    rolling = {"period_labels": {"2026": "TTM Aug 2026"}}
    subs = basis_subs(rolling)

    assert subs[0][0].sub(subs[0][1], "YTD Performance") == "TTM Performance"
    assert basis_subs({"period_labels": {"2026": "Jan 15 – Apr 14 2026"}}) == []
    assert basis_subs({}) == []


def test_without_a_reporting_year_only_dates_move():
    assert _relabel("FY 2025 · 22 July 2026") == "FY 2025 · 28 September 2026"


@pytest.mark.parametrize("latest,expected", [
    ((2024, 12), 2025),       # a complete year plans for the next one
    ((2026, 8), 2026),        # a year still running plans for itself
    (None, 2025),             # no month information → the template's convention
])
def test_the_year_planned_for(latest, expected):
    assert forward_year(2024 if expected == 2025 else 2026, latest) == expected


# ── currency ─────────────────────────────────────────────────────────────────


@pytest.mark.parametrize("authored,expected", [
    ("GWP Performance YoY (€M)", "GWP Performance YoY ({cur}M)"),
    ("€106.5m", "{cur}106.5m"),
    ("£xx,xxxm", "{cur}xx,xxxm"),
    ("Premium grew in Europe", "Premium grew in Europe"),
])
def test_authored_currency_symbols_become_the_reporting_currency(authored, expected):
    assert in_reporting_currency(authored) == expected.format(cur=currency())


# ── one-country portfolio pages ──────────────────────────────────────────────


def test_country_blocks_past_the_selection_come_off_the_table():
    template = _template("product_template.pptx")

    trim = feedback.surplus_country_rows(template, 1)

    first = trim.get("0:2")                 # the first page keeps Country (1) only
    assert first == {"rows": [3, 4]}


def test_nothing_is_trimmed_when_no_country_is_in_scope():
    assert feedback.surplus_country_rows(_template("product_template.pptx"), 0) == {}


# ── the line-of-business icon ────────────────────────────────────────────────


class _Scoped:
    def __init__(self, product):
        self.resolved_filters = {"Product_Line": product}


def test_a_page_about_another_line_loses_the_authored_icon():
    template = _template("product_template.pptx")

    dropped = product_icon.values(template, _Scoped("Casualty"))["drop_shapes"]

    assert "0:29" in dropped                # the ship, and only icons inside the chevron
    assert all(key.endswith(":29") for key in dropped)


def test_the_line_the_page_was_drawn_for_keeps_its_icon():
    assert product_icon.values(_template("product_template.pptx"), _Scoped("Marine")) == {}


# ── end to end: a deck for 2024 says 2024 everywhere ─────────────────────────


def _texts(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def leaves(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from leaves(sh.shapes)
            else:
                yield sh

    for index, slide in enumerate(Presentation(path).slides):
        for sh in leaves(slide.shapes):
            if sh.has_text_frame:
                yield index, sh, sh.text_frame.text
            elif getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        yield index, sh, cell.text
            elif getattr(sh, "has_chart", False) and sh.has_chart:
                chart = sh.chart
                if chart.has_title and chart.chart_title.has_text_frame:
                    yield index, sh, chart.chart_title.text_frame.text
                for series in chart.series:
                    yield index, sh, str(series.name or "")


@pytest.mark.e2e
def test_a_2024_deck_names_2024_on_every_page(tmp_path, monkeypatch):
    from studio.compute import compute_overall
    from studio.template_fill.assemble import assemble_deck

    monkeypatch.setenv("COMMENTARY_MODE", "auto")
    result = compute_overall(filters={"carrier": "Zurich", "country": ["Singapore"],
                                      "year": [2024]})
    out = assemble_deck(result, out_path=str(tmp_path / "deck.pptx"),
                        work_dir=str(tmp_path / "work"))
    texts = list(_texts(out))
    joined = "\n".join(text for _i, _sh, text in texts)

    headers = [t for _i, _sh, t in texts if re.search(r"GWP 20\d\d|20\d\d Rank|SoW% 20\d\d", t)]
    assert headers, "the portfolio and summary tables carry year headers"
    assert all("2024" in h for h in headers), headers
    assert "€" not in joined and "£" not in joined
    assert not re.search(r"^\s*(CY|PY)\s*$", joined, re.M), "periods are named, not CY/PY"
    assert TODAY.strftime("%B") in joined or date.today().strftime("%B") in joined


@pytest.mark.e2e
def test_the_total_bars_plot_the_totals_they_are_labelled_with(tmp_path, monkeypatch):
    from pptx import Presentation

    from studio.compute import compute_overall
    from studio.template_fill.assemble import assemble_deck

    monkeypatch.setenv("COMMENTARY_MODE", "auto")
    result = compute_overall(filters={"carrier": "Zurich", "country": ["Singapore"],
                                      "year": [2024]})
    out = assemble_deck(result, out_path=str(tmp_path / "deck.pptx"),
                        work_dir=str(tmp_path / "work"))
    stacked = [sh for _i, sh, _t in _texts(out)
               if getattr(sh, "has_chart", False) and sh.has_chart
               and "STACKED" in str(sh.chart.chart_type)]
    assert stacked, "the country page's two-bar chart is in the deck"
    bars = [v for v in stacked[0].chart.series[0].values]
    labels = [t for _i, sh, t in _texts(out) if re.fullmatch(r"\s*\S\d+(\.\d)?m\s*", t or "")]
    as_millions = sorted(round(v / 1e6, 1) for v in bars)
    stated = sorted(float(re.sub(r"[^\d.]", "", t)) for t in labels)
    assert as_millions == stated
    assert Presentation(out).slides
