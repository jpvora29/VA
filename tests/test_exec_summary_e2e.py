"""The executive summary end to end, against the real seed book.

    Setup selection -> working book -> premium cube + survey -> evidence -> four stages
    -> verified commentary -> filled template -> audit sidecar -> Generate / canvas doc

The model is either off (``NoModel``: every stage takes its deterministic fallback, the
``STUDIO_AI=off`` path) or an echo model that answers each payload from the facts it was
given — so the model path, the verifier and the deck are exercised without an endpoint.
"""
from __future__ import annotations

import json
import re

import pytest
from pptx import Presentation

from studio.compute import compute_overall
from studio.exec_summary.llm import NoModel
from studio.exec_summary.pipeline import build_exec_summary
from studio.exec_summary.schemas import ModelFinding, StageOutput

SELECTION = {"carrier": "Zurich", "country": ["Singapore", "Hong Kong"], "year": 2025}
PLACEHOLDER = re.compile(r"\bx{2,}|#XX|\$X+M|XX\b|#N", re.I)


@pytest.fixture(autouse=True)
def _offline(monkeypatch):
    monkeypatch.setenv("STUDIO_COMMENTARY_CACHE", "off")


class EchoModel:
    """Answers every stage by quoting the first facts in its payload, cited correctly."""

    label = "echo"

    def __init__(self):
        self.stages = []

    def complete(self, stage, system, user):
        self.stages.append(stage)
        data = json.loads(user.split("INPUT JSON:\n", 1)[1].rsplit("\n\nReturn", 1)[0])
        facts = re.findall(r'"(carrier_premium|marsh_premium)":"(e\d+) \| ([^"]+)"', user)
        findings = [ModelFinding(rule_id=data["allowed_rule_ids"][0], severity="context",
                                 statement=f"The {'carrier' if m == 'carrier_premium' else 'Marsh'} "
                                           f"figure here is {display}.",
                                 interpretation="", recommendation="", evidence_ids=[eid],
                                 qualifications=[])
                    for m, eid, display in facts[:2]]
        return StageOutput(scope_id=data["scope_id"], findings=findings, priority_refs=[],
                           missing_references=[])


def _result(**kw):
    return compute_overall(filters=SELECTION, **kw)


def _deck_text(path) -> list:
    prs = Presentation(path)
    return [" ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
            for slide in prs.slides]


def test_ai_off_builds_a_complete_deck_from_deterministic_findings(tmp_path):
    out = tmp_path / "exec.pptx"
    built = build_exec_summary(_result(), out_path=str(out), data_basis="premium_survey",
                               model=NoModel())
    prs = Presentation(built.deck_path)
    assert len(prs.slides) == 4
    cover, summary, detail, _ = _deck_text(built.deck_path)
    assert "Zurich Executive Summary" in cover
    assert "Zurich — Carrier Performance Executive Summary (FY2025)" in summary
    assert "Overall, Zurich's Marsh-placed premium" in summary and "survey score" in summary
    assert summary.index("In Singapore,") < summary.index("In Hong Kong,") < summary.index(
        "By product,")
    for text in (cover, summary, detail):
        assert not PLACEHOLDER.search(text), PLACEHOLDER.search(text)

    tables = [s for s in prs.slides[2].shapes if s.has_table]
    lob = next(t.table for t in tables if len(t.table.columns) >= 7)
    header = [lob.cell(0, c).text for c in range(len(lob.columns))]
    assert header == ["Line of Business", "Marsh GWP", "Carrier GWP", "YoY",
                      "Share of Portfolio", "Share of Wallet", "Rank", "Peers GWP"]
    assert len(lob.rows) - 1 == len(built.analysis.product_totals) == 6
    mixes = [float(lob.cell(r, 4).text.rstrip("%")) for r in range(1, len(lob.rows))]
    assert sum(mixes) == pytest.approx(100.0, abs=0.5)          # shares of the carrier's own book

    rank = next(t.table for t in tables if len(t.table.columns) == 2)
    assert [rank.cell(r, 0).text for r in range(1, len(rank.rows))] == ["Singapore", "Hong Kong"]
    assert all(re.fullmatch(r"#\d+ \((?:[+-]\d+|=|new)\)", rank.cell(r, 1).text)
               for r in range(1, len(rank.rows)))

    charts = sorted((s for s in prs.slides[2].shapes if getattr(s, "has_chart", False) and s.has_chart),
                    key=lambda s: s.left)
    premium = charts[0].chart.plots[0]
    assert list(premium.categories) == ["Singapore", "Hong Kong"]
    sg = built.analysis.countries[0].metrics
    assert premium.series[0].values[0] == pytest.approx(round(sg.carrier_cur / 1e6, 1))
    assert premium.series[1].values[0] == pytest.approx(round(sg.peer.mean_cur / 1e6, 1))
    assert "Top 3 lines hold" in detail


def test_every_summary_sentence_traces_to_cited_evidence(tmp_path):
    built = build_exec_summary(_result(), out_path=str(tmp_path / "e.pptx"),
                               data_basis="premium_survey", model=NoModel())
    envelope = json.loads(open(built.report_path, encoding="utf-8").read())
    evidence = {e["evidence_id"]: e for e in envelope["evidence"]}
    for finding in envelope["overall"]["findings"]:
        assert finding["validation_status"] == "verified"
        assert finding["evidence_ids"]
        displays = " ".join(evidence[i]["display"] for i in finding["evidence_ids"])
        for number in re.findall(r"\d+\.\d", finding["statement"]):
            assert number in displays, (number, finding["statement"])
    assert envelope["coverage_summary"]["stages_fallback"] == envelope["coverage_summary"]["stages"]
    assert envelope["opportunity_register"], "the full register is kept for review"


def test_the_model_path_reaches_the_slide_only_after_verification(tmp_path):
    model = EchoModel()
    built = build_exec_summary(_result(), out_path=str(tmp_path / "e.pptx"),
                               data_basis="premium_survey", model=model)
    assert {"product", "opportunity", "country", "overall"} <= set(model.stages)
    overall = built.commentary.overall
    assert overall.source == "model"
    summary = _deck_text(built.deck_path)[1]
    assert overall.findings[0].statement in summary
    # Call count = products + opportunities + countries + 1 (design Section 18).
    children = built.commentary.children
    assert len(model.stages) == len(children) + len(built.commentary.countries) + 1


def test_country_specific_peer_pins_stay_country_specific(tmp_path):
    result = _result(peers_by_country={"Singapore": ["AIG", "Chubb", "Allianz", "AXA XL", "QBE"],
                                       "Hong Kong": ["Sompo", "Tokio Marine", "MS&AD",
                                                     "Swiss Re", "Berkshire"]})
    built = build_exec_summary(result, out_path=str(tmp_path / "e.pptx"),
                               data_basis="premium", model=NoModel())
    envelope = json.loads(open(built.report_path, encoding="utf-8").read())
    peers = envelope["peer_versions_by_country"]
    assert peers["Singapore"]["members"] != peers["Hong Kong"]["members"]
    assert peers["Singapore"]["source"] == "setup_pin"
    assert "QBE" in peers["Singapore"]["members"] and "QBE" not in peers["Hong Kong"]["members"]
    # Named peers never reach the deck.
    text = " ".join(_deck_text(built.deck_path)[:3])
    for name in ("AIG", "Chubb", "Sompo", "Tokio Marine"):
        assert not re.search(rf"\b{re.escape(name)}\b", text)


def test_a_premium_only_run_leaves_the_survey_out_and_says_so(tmp_path):
    built = build_exec_summary(_result(), out_path=str(tmp_path / "e.pptx"),
                               data_basis="premium", model=NoModel())
    summary = _deck_text(built.deck_path)[1]
    assert "survey score" not in summary
    assert "Survey: not included" in summary
    assert built.analysis.overall_survey is None


def test_generate_builds_the_executive_summary_for_the_canvas(monkeypatch):
    from studio.authoring import generate as G

    monkeypatch.setenv("STUDIO_AI", "off")
    monkeypatch.setenv("COMMENTARY_MODE", "auto")
    monkeypatch.setattr(G, "ensure_rendered_slide_backgrounds", lambda path, n, only=None: [None] * n)
    monkeypatch.setattr(G, "prune_cache", lambda paths: None)
    G._exec_summary_for.cache_clear()
    phases = []
    selection = {"report": "exec_summary", "filters": SELECTION, "data_basis": "premium"}
    docs = G.build_documents(selection, lambda phase, msg: phases.append(phase))
    assert docs.doc is None
    assert docs.tdoc["report"] == "exec_summary" and docs.tdoc["assembled"] is True
    assert docs.tdoc["n_slides"] == 4 and docs.tdoc["template_path"].endswith("_Executive_Summary.pptx")
    assert phases == ["data", "assemble", "render"]


def test_setup_offers_the_deliverable_and_the_bar_names_it():
    from studio.authoring.setup_summary import bar_state
    from studio.page.authoring import deliverable as DV
    from studio.page.authoring.setup_layout import compose_setup

    assert "id='studio-deliverable'" in str(compose_setup())
    _, _, summary, ready = bar_state("Zurich", 4, "premium", "carrier_leadership", "governed",
                                     DV.REPORT_EXEC_SUMMARY)
    assert ready and summary.startswith("Executive summary · 4 pages")
    assert DV.page_count(DV.REPORT_QBR) is None


def test_the_progress_card_names_the_executive_summary():
    from studio.authoring.jobs import _shape_label

    assert _shape_label({"report": "exec_summary"}) == "Executive summary"
    assert _shape_label({"report": "qbr"}) == "Full QBR"
