"""The executive summary reads as a story, names the real period, and fits the slide.

* the period is named for what the data covers (H1, Q1, 9M …) and compared like for like;
* the summary names WHAT moved — the leading product, its rank move, the industry inside
  it, the product that held the total back — rather than listing figures;
* the line-of-business table carries the top 7 lines by carrier premium.

Synthetic where the point is arithmetic; the seed book (with its second half-year cut
off) where the point is the real flow.
"""
from __future__ import annotations

import re
import shutil
import sqlite3

import pandas as pd
import pytest
from sqlalchemy import create_engine

from studio.exec_summary.analysis import analyse
from studio.exec_summary.config import load_config
from studio.exec_summary.context import AnalysisContext
from studio.exec_summary.coverage import month_number, period_label
from studio.exec_summary.fallback import overall_fallback
from studio.exec_summary.slides import MAX_TABLE_ROWS, table_rows
from studio.exec_summary.verify import _direction_issue
from studio.exec_summary.evidence import EvidenceStore, UNIT_MONEY

CFG = load_config()
M = 1_000_000.0


# ── the period (H1, Q1, 9M …) ────────────────────────────────────────────────


@pytest.mark.parametrize("months, label", [
    (range(1, 13), "FY2026"), (range(1, 7), "H1 2026"), (range(7, 13), "H2 2026"),
    (range(1, 4), "Q1 2026"), (range(4, 7), "Q2 2026"), (range(1, 10), "9M 2026"),
    (range(1, 6), "YTD May 2026"), (range(2, 5), "Feb–Apr 2026"),
])
def test_the_period_is_named_for_the_months_the_data_covers(months, label):
    assert period_label(list(months), 2026) == label


@pytest.mark.parametrize("raw, n", [("January", 1), ("jun", 6), ("Sept", 9), ("07", 7),
                                    ("2026-03", 3), ("2026-11-15", 11), ("Q1", None)])
def test_month_values_are_read_in_any_common_spelling(raw, n):
    assert month_number(raw) == n


@pytest.fixture(scope="module")
def h1_engine(tmp_path_factory):
    """The seed book with July–December of 2025 removed: a book loaded to June."""
    from studio.data import get_engine

    path = tmp_path_factory.mktemp("h1") / "h1.db"
    shutil.copyfile(get_engine().url.database, path)
    with sqlite3.connect(path) as conn:
        conn.execute("DELETE FROM GPR WHERE Year = 2025 AND Month_Name IN "
                     "('July','August','September','October','November','December')")
    engine = create_engine(f"sqlite:///{path}")
    yield engine, str(path)
    engine.dispose()


def test_a_half_year_book_is_named_h1_and_compared_with_the_prior_h1(h1_engine, tmp_path,
                                                                    monkeypatch):
    from studio.compute import compute_overall
    from studio.exec_summary.llm import NoModel
    from studio.exec_summary.pipeline import build_exec_summary

    monkeypatch.setenv("STUDIO_COMMENTARY_CACHE", "off")
    engine, path = h1_engine
    result = compute_overall(filters={"carrier": "Zurich", "country": ["Singapore", "Hong Kong"],
                                      "year": 2025}, engine=engine)
    built = build_exec_summary(result, out_path=str(tmp_path / "h1.pptx"), data_basis="premium",
                               model=NoModel())
    ctx, overall = built.analysis.context, built.analysis.overall
    assert (ctx.current_label, ctx.prior_label) == ("H1 2025", "H1 2024")

    h1 = ("January", "February", "March", "April", "May", "June")
    with sqlite3.connect(path) as conn:
        by_year = dict(conn.execute(
            f"SELECT Year, SUM(Premium) FROM GPR WHERE Carrier_Group = 'Zurich' AND Country IN "
            f"('Singapore','Hong Kong') AND Month_Name IN ({','.join('?' * 6)}) "
            f"AND Year IN (2024, 2025) GROUP BY Year", h1).fetchall())
    assert overall.carrier_cur == pytest.approx(by_year[2025])
    assert overall.carrier_prior == pytest.approx(by_year[2024])     # H1 2024, not FY2024

    from pptx import Presentation
    summary = " ".join(s.text_frame.text for s in Presentation(built.deck_path).slides[1].shapes
                       if s.has_text_frame)
    assert "Executive Summary (H1 2025)" in summary
    assert "in H1 2025" in summary and "FY2025" not in summary


# ── the story: named movers, not a list of figures ───────────────────────────


def _book(products=("Cyber", "Property", "Casualty")) -> pd.DataFrame:
    rows = []
    plan = {"Cyber": (10, 30, "Technology"), "Property": (40, 42, "Manufacturing"),
            "Casualty": (30, 24, "Wholesale")}
    for p in products:
        prior, cur, industry = plan.get(p, (5, 6, "Retail"))
        for year, focal in ((2024, prior), (2025, cur)):
            rows += [
                {"country": "SG", "product": p, "industry": industry, "segment": "Corporate",
                 "carrier": "Focal", "year": year, "premium": focal * M},
                {"country": "SG", "product": p, "industry": "Other", "segment": "Corporate",
                 "carrier": "Focal", "year": year, "premium": 1 * M},
                {"country": "SG", "product": p, "industry": industry, "segment": "Corporate",
                 "carrier": "Rival", "year": year, "premium": (100 if year == 2024 else 120) * M},
            ]
    return pd.DataFrame(rows)


def _analysis(products=("Cyber", "Property", "Casualty")):
    ctx = AnalysisContext(carrier="Focal", countries=("SG",), current_year=2025, prior_year=2024,
                          current_label="H1 2025", prior_label="H1 2024", currency="USD",
                          symbol="$")
    return analyse(ctx, _book(products), {}, None, CFG)


def test_the_summary_reads_overall_then_country_then_products_gaps_and_openings():
    a = _analysis()
    findings = overall_fallback(a, {}).findings
    rules = [f.rule_id for f in findings]
    statements = [f.statement for f in findings]
    # The order the user asked for (2026-10-17).
    assert rules[:2] == ["EXEC_OVERALL", "EXEC_COUNTRY"]
    assert rules.index("EXEC_PRODUCT") > rules.index("EXEC_COUNTRY")
    assert rules.index("EXEC_GAP") > rules.index("EXEC_PRODUCT")
    # 1. Overall: carrier premium + YoY and the Marsh book — no product, no country.
    overall = statements[0]
    assert overall.startswith("Overall, Focal's Marsh-placed premium grew 19.3% to $99.0M "
                              "in H1 2025")
    assert "the Marsh book grew 19.8% to $459.0M" in overall
    assert not any(n in overall for n in ("Cyber", "Property", "Casualty", "SG"))
    # 2. The country: share of wallet and rank.
    assert statements[1].startswith("In SG,") and "share of wallet at" in statements[1]
    assert re.search(r"at rank #\d", statements[1])
    # 3. Products: what added most (rank move, industry inside), what held it back.
    product = [s for f, s in zip(findings, statements) if f.rule_id == "EXEC_PRODUCT"]
    assert product[0].startswith("By product, Cyber added the most, $20.0M")
    assert "mostly in Technology" in product[0]
    assert re.search(r"(climbing to|slipping to|holding) #\d", product[0])
    assert product[1].startswith("Casualty held the total back: it fell 19.4%")
    assert "mostly in Wholesale" in product[1]
    # 4. Gaps: the product behind its Marsh book.
    gap = statements[rules.index("EXEC_GAP")]
    assert "Casualty trails its Marsh book's growth by" in gap


def test_a_multi_country_summary_says_where_each_product_moved():
    rows = []
    for country, cyber in (("SG", (10, 30)), ("HK", (10, 14))):
        for year, focal in ((2024, cyber[0]), (2025, cyber[1])):
            rows += [
                {"country": country, "product": "Cyber", "industry": "Tech", "segment": "C",
                 "carrier": "Focal", "year": year, "premium": focal * M},
                {"country": country, "product": "Cyber", "industry": "Tech", "segment": "C",
                 "carrier": "Rival", "year": year, "premium": 50 * M},
            ]
    ctx = AnalysisContext(carrier="Focal", countries=("SG", "HK"), current_year=2025,
                          prior_year=2024, current_label="FY2025", prior_label="FY2024",
                          currency="USD", symbol="$")
    a = analyse(ctx, pd.DataFrame(rows), {}, None, CFG)
    findings = overall_fallback(a, {}).findings
    assert [f.rule_id for f in findings][:3] == ["EXEC_OVERALL", "EXEC_COUNTRY", "EXEC_COUNTRY"]
    assert findings[1].statement.startswith("In SG,")              # largest market first
    assert findings[2].statement.startswith("In HK,")
    product = next(f.statement for f in findings if f.rule_id == "EXEC_PRODUCT")
    assert "($20.0M in SG and $4.0M in HK)" in product


def test_the_verifier_holds_the_overall_structure():
    """A model answer that opens with a product, or omits WHERE, is refused — and verified
    findings are put back in section order whatever order they came in."""
    from studio.exec_summary import payloads as P
    from studio.exec_summary.schemas import ModelFinding, StageOutput
    from studio.exec_summary.verify import names_universe, verify_stage

    a = _analysis()
    payload = P.overall_payload(a, {}, {})
    w = a.overall_evidence
    disp = a.evidence.get(w["carrier_premium"]).display

    def finding(rule, text, ids):
        return ModelFinding(rule_id=rule, severity="context", statement=text,
                            interpretation="", recommendation="", evidence_ids=ids,
                            qualifications=[])

    out = StageOutput(scope_id="ALL/overall", priority_refs=[], missing_references=[], findings=[
        finding("EXEC_GAP", f"The gap is share at {disp}.", [w["carrier_premium"]]),
        finding("EXEC_OVERALL", f"Cyber led premium to {disp}.", [w["carrier_premium"]]),
        finding("EXEC_OVERALL", f"Overall premium reached {disp}.", [w["carrier_premium"]]),
    ])
    verified = verify_stage(out, payload, a.evidence, universe=names_universe(_book()))
    assert [f.rule_id for f in verified.findings] == ["EXEC_OVERALL", "EXEC_GAP"]
    assert any("overall sentence names 'Cyber'" in e for e in verified.errors)


def test_every_fallback_sentence_still_passes_the_claim_checks():
    from studio.exec_summary import payloads as P
    from studio.exec_summary.stages import StageCall, StageEnv, run_stage
    from studio.exec_summary.llm import NoModel
    from functools import partial
    from studio.exec_summary import fallback as FB

    a = _analysis()
    env = StageEnv(model=NoModel(), store=a.evidence, run_id="t", rule_version="r",
                   budgets=CFG.budgets, concurrency=1)
    result = run_stage(StageCall(P.overall_payload(a, {}, {}), partial(FB.overall_fallback, a, {})),
                       env)
    assert result.errors == ("model unavailable (none)",)
    assert len(result.findings) >= 3


def test_a_repeated_magnitude_is_not_read_against_the_wrong_verb():
    store = EvidenceStore()
    gain = store.add("carrier_premium_change", "s", "A", "p", 0.5 * M, "+$0.5M", UNIT_MONEY)
    loss = store.add("carrier_premium_change", "s", "B", "p", -0.5 * M, "-$0.5M", UNIT_MONEY)
    ok = "Inside Casualty, Retail added $0.5M while Wholesale lost $0.5M."
    assert _direction_issue(ok, store, [gain, loss]) == ""
    wrong = "Inside Casualty, Wholesale grew $0.5M."
    assert "rose beside a negative" in _direction_issue(wrong, store, [loss])


# ── the detail table fits ────────────────────────────────────────────────────


def test_the_table_carries_the_top_seven_lines_by_carrier_premium():
    products = ("Cyber", "Property", "Casualty", "Marine", "Energy", "Surety", "Aviation",
                "Credit", "Terror")
    a = _analysis(products)
    rows = table_rows(a)
    assert MAX_TABLE_ROWS == 7 and len(rows) == 7
    premiums = [p.carrier_cur for p in a.product_totals]
    assert premiums == sorted(premiums, reverse=True)
    assert [r.product for r in rows] == [p.scope.product for p in a.product_totals[:7]]
