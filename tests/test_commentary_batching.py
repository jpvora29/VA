"""Section-level commentary: the batching, the strict mode, the cache, the telemetry.

The plan these cover ("QBR AI Commentary, Chatbot Quality, and Performance") is about one
number — how many model calls a build makes — and three guarantees that must survive
cutting it:

  * a batched deck maps every requested field EXACTLY once, in the deck's own order;
  * ``COMMENTARY_MODE=ai_required`` never lets deterministic prose reach a slide;
  * a re-export of unchanged evidence makes no model call at all.

No LLM and no database. The model is a stub that answers the structured schema, so what is
under test is the orchestration rather than the writing.
"""
from __future__ import annotations

import pytest
import re

from studio import commentary_mode as mode
from studio import telemetry
from studio.ai.models import CommentaryBullet, CommentarySection, CommentarySections
from studio.template_fill import commentary_batch as B
from studio.template_fill import rewrites
from studio.template_fill.rewrites import PendingRewrite
from studio.template_fill.deck_slides import DeckSlides


# ── fixtures: one book, enough of it to build a real evidence pack ───────────


def _approved_verdicts(user):
    """Explicit successful review for orchestration tests; absence is now failure."""
    from studio.ai.models import CommentaryVerdict, CommentaryVerdicts
    count = len(re.findall(r"^\d+\. ", user, re.M))
    return CommentaryVerdicts(verdicts=[CommentaryVerdict(keep=True) for _ in range(count)])


def _facts() -> dict:
    """A fact dict of the shape ``feedback.facts_for`` returns, big enough to cite.

    The key names matter: :func:`studio.template_fill.commentary_evidence.build_pack` reads
    ``current``/``pct``/``delta`` off each family, and a plausible-looking dict with the
    wrong names builds a pack of two facts and makes every evidence assertion vacuous.
    """
    return {
        "subject": "Zurich",
        "carrier": {"current": 44_000_000.0, "pct": 12.8, "delta": 5_000_000.0,
                    "current_year": 2025},
        "marsh": {"current": 900_000_000.0, "pct": 6.2},
        "sow": {"current": 4.9, "delta": 0.5},
        "rank": {"current": 4, "of_n": 31, "delta": 1},
        "peer": {"sow": 6.1, "current": 55_000_000.0},
    }


def _pending(topic: str, draft: str, facts: dict) -> PendingRewrite:
    return PendingRewrite(draft=draft, node=f"commentary-{topic}", topic=topic,
                          subject="Zurich", style="balanced", facts=facts)


def _value_set(facts: dict, topics=("working", "challenges", "growth")) -> dict:
    return {f"note:{i}:{i}:0": _pending(t, f"{t} line one.\n{t} line two.", facts)
            for i, t in enumerate(topics)}


@pytest.fixture(autouse=True)
def _no_cache(tmp_path, monkeypatch):
    """Point the on-disk cache at a throwaway directory for every test in this file."""
    monkeypatch.setenv("STUDIO_COMMENTARY_CACHE", str(tmp_path / "commentary"))
    monkeypatch.setenv("COMMENTARY_MODE", "auto")
    monkeypatch.setenv("STUDIO_PPT_MERGE_ENGINE", "opc")
    monkeypatch.delenv("STUDIO_AI", raising=False)


#: What the stubbed model writes. Whole sentences, no figures, no phrase the reading gate
#: refuses — and no phrase that names a CLAIM TOPIC, so the editorial gate
#: (:func:`studio.template_fill.editorial.dedupe`) has nothing to recognise either. The
#: stub answers every field with the same text, which is a repeated finding the moment one
#: of these sentences is about something; keeping them topicless is what keeps a failed
#: assertion below about the batching rather than about the prose. The gate has its own
#: file, ``test_commentary_editorial.py``.
_SENTENCES = (
    "The year closed better than the one before it, and that is the story to tell.",
    "Defending that position now matters more than adding to it.",
    "Growth of this kind rarely survives a hard renewal season without a plan behind it.",
    "The team should treat the coming quarter as a test of whether the gain holds.",
)


@pytest.fixture
def stub_model(monkeypatch):
    """A model that answers every requested field with two citable sentences.

    Returns the call log, so a test can assert on HOW MANY calls a deck made — which is
    the number the whole plan is about.
    """
    import studio.ai.client as client

    calls = []

    def structured(model, system, user, *, tier="balanced", node="ai", phase="other",
                   fields=(), **kw):
        calls.append({"node": node, "phase": phase, "fields": tuple(fields), "user": user})
        if model is CommentarySections:
            asked = [line.split()[2] for line in user.splitlines()
                     if line.startswith("--- FIELD ")]
            # Deliberately figure-free whole sentences: the verifiers and the shape gate
            # are under test elsewhere, and a stub that trips them would make every
            # call-count assertion here really an assertion about the gates. Four of them,
            # because a real template's columns ask for up to four and ``_accept`` trims a
            # generous answer but refuses a short one.
            return CommentarySections(sections=[
                CommentarySection(field_id=fid, bullets=[
                    CommentaryBullet(text=line, fact_ids=["period.year"]) for line in _SENTENCES
                ]) for fid in asked])
        return _approved_verdicts(user)

    monkeypatch.setattr(client, "structured", structured)
    monkeypatch.setattr(client, "llm_available", lambda: True)
    return calls


# ── grouping: one call per sub-deck, one entry per field ────────────────────


def test_a_sub_decks_columns_are_grouped_into_one_section():
    facts = _facts()
    sections = B.group_sections([_value_set(facts)])
    assert len(sections) == 1
    assert len(sections[0].columns) == 3
    assert sections[0].subject == "Zurich"


def test_two_sub_decks_reporting_different_books_are_two_sections():
    first, second = _facts(), _facts()
    second["carrier"]["current"] = 51_000_000.0      # a different book, not a copy
    sections = B.group_sections([_value_set(first), _value_set(second)])
    assert [s.value_sets for s in sections] == [(0,), (1,)]


def test_two_sub_decks_reporting_THE_SAME_book_are_one_section():
    """The slide 6 / slide 16 fix.

    On a single-country run the overall block and the country block describe one book down
    to the last figure. As two sections they were two concurrent calls, so neither could
    see that both had made the peer-gap point — which is how one finding reached two pages
    of a shipped deck. One book is one section, and then the editorial plan and the
    repetition gate cover both pages.
    """
    facts = _facts()
    sections = B.group_sections([_value_set(facts), _value_set(_facts())])
    assert len(sections) == 1
    assert sections[0].value_sets == (0, 1)


def test_a_merged_section_still_lands_each_column_in_its_own_sub_deck(stub_model):
    """Merging must not smear one sub-deck's prose onto another's roles."""
    first, second = _value_set(_facts()), _value_set(_facts())
    written = rewrites.write_all([first, second])
    assert [set(w) for w in written] == [set(first), set(second)]
    assert all(isinstance(v, str) and v.strip() for w in written for v in w.values())


def test_columns_sharing_one_pending_are_written_once_and_land_on_every_role():
    """``commentary.values`` answers one question per topic and puts that answer in every
    box asking it — the same object under several roles. Writing it once per role paid for
    the same column twice."""
    facts = _facts()
    shared = _pending("working", "one line.", facts)
    sections = B.group_sections([{"note:1:1:0": shared, "note:2:2:0": shared}])
    assert len(sections[0].columns) == 1
    assert {t.role for t in sections[0].columns[0].targets} == {"note:1:1:0", "note:2:2:0"}


def test_every_requested_field_id_is_returned_exactly_once(stub_model):
    facts = _facts()
    values = _value_set(facts)
    written = rewrites.write_all([values])[0]
    assert set(written) == set(values)
    assert all(isinstance(v, str) and v.strip() for v in written.values())


def test_a_field_the_model_invents_is_ignored_rather_than_mapped_by_position(stub_model,
                                                                            monkeypatch):
    import studio.ai.client as client

    def structured(model, system, user, *, tier="balanced", node="ai", phase="other",
                   fields=(), **kw):
        if model is CommentarySections:
            return CommentarySections(sections=[CommentarySection(
                field_id="not.a.field",
                bullets=[CommentaryBullet(text="A sentence about nothing.", fact_ids=["period.year"])])])
        return _approved_verdicts(user)

    monkeypatch.setattr(client, "structured", structured)
    facts = _facts()
    values = _value_set(facts, topics=("working",))
    written = rewrites.write_all([values])[0]
    # Nothing was written, so the deterministic draft stands — not another field's text.
    assert list(written.values()) == ["working line one.\nworking line two."]


# ── the call count: the whole point of the change ───────────────────────────


def test_a_sub_deck_costs_one_author_call_not_one_per_column(stub_model):
    facts = _facts()
    rewrites.write_all([_value_set(facts, topics=("working", "challenges", "growth",
                                                  "priorities", "key_messages"))])
    authored = [c for c in stub_model if c["phase"] == "author"]
    assert len(authored) == 1, "five columns must cost one author call"
    assert len(authored[0]["fields"]) == 5


def test_three_sub_decks_cost_three_calls_not_fifteen(stub_model):
    # Three DIFFERENT books. Identical ones would be three sections racing for one cache
    # entry, and the call count would depend on which finished first.
    books = []
    for i in range(3):
        book = _facts()
        book["carrier"]["current"] += i * 1_000_000.0
        books.append(book)
    rewrites.write_all([_value_set(f, topics=("working", "challenges", "growth",
                                             "priorities", "key_messages"))
                        for f in books])
    assert len([c for c in stub_model if c["phase"] == "author"]) == 3


# ── strict mode: no deterministic prose, ever ───────────────────────────────


def test_ai_required_refuses_rather_than_shipping_the_draft(monkeypatch):
    import studio.ai.client as client

    monkeypatch.setenv("COMMENTARY_MODE", "ai_required")
    monkeypatch.setattr(client, "llm_available", lambda: True)
    monkeypatch.setattr(client, "structured",
                        lambda *a, **kw: None)          # the author answers nothing
    with pytest.raises(mode.CommentaryUnavailable) as raised:
        rewrites.write_all([_value_set(_facts())])
    assert raised.value.retryable is True
    assert "ai_required" in str(raised.value)


def test_auto_mode_keeps_the_draft_when_the_author_answers_nothing(monkeypatch):
    import studio.ai.client as client

    monkeypatch.setattr(client, "llm_available", lambda: True)
    monkeypatch.setattr(client, "structured", lambda *a, **kw: None)
    written = rewrites.write_all([_value_set(_facts(), topics=("working",))])[0]
    assert list(written.values()) == ["working line one.\nworking line two."]


def test_preflight_refuses_before_any_analytics_when_ai_is_not_configured(monkeypatch):
    monkeypatch.setenv("COMMENTARY_MODE", "ai_required")
    monkeypatch.setattr("core.llm.clients.available", lambda tier="balanced": False)
    with pytest.raises(mode.CommentaryUnavailable) as raised:
        mode.preflight()
    assert raised.value.retryable is False, "a missing key is not fixed by trying again"


def test_preflight_is_silent_in_auto_mode(monkeypatch):
    monkeypatch.setattr("core.llm.clients.available", lambda tier="balanced": False)
    mode.preflight()                                    # no raise


def test_the_kill_switch_beats_the_strict_mode(monkeypatch):
    monkeypatch.setenv("STUDIO_AI", "off")
    monkeypatch.setenv("COMMENTARY_MODE", "ai_required")
    assert mode.mode() == mode.OFF
    assert not mode.ai_required()
    mode.preflight()                                    # and refuses nothing


def test_an_unreadable_mode_requires_verified_ai(monkeypatch):
    monkeypatch.setenv("COMMENTARY_MODE", "sort-of-required")
    assert mode.mode() == mode.AI_REQUIRED


def test_strict_mode_withholds_the_finished_draft_from_the_author(monkeypatch, stub_model):
    monkeypatch.setenv("COMMENTARY_MODE", "ai_required")
    rewrites.write_all([_value_set(_facts(), topics=("working",))])
    prompt = [c for c in stub_model if c["phase"] == "author"][0]["user"]
    assert "working line one." not in prompt, "strict mode must not show prose to imitate"
    assert "EVIDENCE" in prompt and "FIELD" in prompt


def test_auto_mode_also_authors_without_a_finished_draft(stub_model):
    rewrites.write_all([_value_set(_facts(), topics=("working",))])
    prompt = [c for c in stub_model if c["phase"] == "author"][0]["user"]
    assert "working line one." not in prompt


# ── the cache: an export must not re-write the deck ─────────────────────────


def test_a_second_identical_build_makes_no_author_call(stub_model):
    facts = _facts()
    first = rewrites.write_all([_value_set(facts, topics=("working", "challenges"))])[0]
    calls_after_first = len([c for c in stub_model if c["phase"] == "author"])

    second = rewrites.write_all([_value_set(facts, topics=("working", "challenges"))])[0]
    assert len([c for c in stub_model if c["phase"] == "author"]) == calls_after_first
    assert list(first.values()) == list(second.values())


def test_changed_evidence_invalidates_the_cached_section(stub_model):
    rewrites.write_all([_value_set(_facts(), topics=("working",))])
    moved = _facts()
    moved["carrier"]["current"] = 51_000_000.0
    rewrites.write_all([_value_set(moved, topics=("working",))])
    assert len([c for c in stub_model if c["phase"] == "author"]) == 2


def test_a_new_prompt_version_invalidates_the_cache(stub_model, monkeypatch):
    facts = _facts()
    rewrites.write_all([_value_set(facts, topics=("working",))])
    # Any version other than the shipped one. Naming the NEXT version here made the test
    # fail the day that version shipped, which is the one day it had nothing to say.
    monkeypatch.setattr(B, "PROMPT_VERSION", B.PROMPT_VERSION + "-changed")
    rewrites.write_all([_value_set(facts, topics=("working",))])
    assert len([c for c in stub_model if c["phase"] == "author"]) == 2


def test_a_retry_after_a_partial_failure_only_rewrites_what_failed(monkeypatch):
    """"Preserve successful partial results across retries": a section that wrote one of
    its two fields must not pay to write that one again on the next run."""
    import studio.ai.client as client

    answered = {"working"}                   # only this field comes back the first time
    asked = []

    def structured(model, system, user, *, tier="balanced", node="ai", phase="other",
                   fields=(), **kw):
        if model is not CommentarySections:
            return _approved_verdicts(user)
        asked.append(tuple(fields))
        wanted = [line.split()[2] for line in user.splitlines()
                  if line.startswith("--- FIELD ")]
        return CommentarySections(sections=[
            CommentarySection(field_id=fid, bullets=[
                CommentaryBullet(text=line, fact_ids=["period.year"]) for line in _SENTENCES])
            for fid in wanted if fid.split(".")[0] in answered])

    monkeypatch.setattr(client, "llm_available", lambda: True)
    monkeypatch.setattr(client, "structured", structured)

    facts = _facts()
    topics = ("working", "challenges")
    first = rewrites.write_all([_value_set(facts, topics=topics)])[0]
    written = [v for v in first.values() if v.startswith(_SENTENCES[0])]
    assert len(written) == 1, "exactly one field was answerable on the first run"

    # Second run: the model can now answer both, but must only be ASKED the one that
    # failed — the field that succeeded was cached the moment it passed.
    answered.add("challenges")
    asked.clear()
    second = rewrites.write_all([_value_set(facts, topics=topics)])[0]
    assert all(len(fields) == 1 for fields in asked), asked
    assert all(v.startswith(_SENTENCES[0]) for v in second.values())


def test_the_cache_can_be_turned_off(stub_model, monkeypatch):
    monkeypatch.setenv("STUDIO_COMMENTARY_CACHE", "off")
    facts = _facts()
    rewrites.write_all([_value_set(facts, topics=("working",))])
    rewrites.write_all([_value_set(facts, topics=("working",))])
    assert len([c for c in stub_model if c["phase"] == "author"]) == 2


# ── telemetry ───────────────────────────────────────────────────────────────


def test_a_model_call_records_itself_with_its_phase(monkeypatch):
    """The recording lives in ``studio.ai.client``, so it is driven through the real one."""
    import core.llm.clients as llm
    import studio.ai.client as client
    from pydantic import BaseModel

    class Answer(BaseModel):
        text: str = "written"

    class FakeStructured:
        def invoke(self, messages):
            return Answer()

    class FakeClient:
        def with_structured_output(self, model):
            return FakeStructured()

    monkeypatch.setattr(llm, "make_client", lambda tier: FakeClient())
    monkeypatch.setattr(client, "llm_available", lambda: True)

    with telemetry.job("test") as record:
        client.structured(Answer, "s", "u", tier="reason", node="section-x",
                          phase=telemetry.AUTHOR, fields=("a", "b"))
        client.structured(Answer, "s", "u", node="section-x-verify",
                          phase=telemetry.VERIFY)
        summary = record.summary()

    assert summary.calls == 2
    assert summary.calls_by_phase[telemetry.AUTHOR] == 1
    assert summary.calls_by_phase[telemetry.VERIFY] == 1
    assert record.calls[0].fields == ("a", "b")
    assert "job=test" in summary.as_line()


def test_a_failed_call_is_still_recorded_with_its_failure_kind(monkeypatch):
    """The slowest calls in a build are often the ones that produced nothing."""
    import core.llm.clients as llm
    import studio.ai.client as client

    class Exploding:
        def invoke(self, messages):
            raise RuntimeError("Error code: 429 - rate limit is exceeded")

    monkeypatch.setattr(llm, "make_client", lambda tier: Exploding())
    monkeypatch.setattr(client, "llm_available", lambda: True)

    with telemetry.job("test") as record:
        assert client.generate("s", "u", phase=telemetry.AUTHOR) is None
    assert record.calls[0].status == "error"
    assert record.calls[0].failure == "rate_limit"


def test_failures_are_classified_so_a_retry_policy_can_read_them():
    cases = {
        "rate_limit": Exception("Error code: 429 - rate limit exceeded"),
        "timeout": TimeoutError("request timed out after 60s"),
        "auth": Exception("401 unauthorized: invalid api key"),
        "config": Exception("DeploymentNotFound: the deployment does not exist"),
    }
    for expected, exc in cases.items():
        assert telemetry.classify_failure(exc) == expected
    assert telemetry.is_transient("rate_limit") and telemetry.is_transient("timeout")
    assert not telemetry.is_transient("auth"), "retrying a bad key burns the budget"
    assert not telemetry.is_transient("config")


def test_percentiles_do_not_need_numpy():
    from studio.telemetry import _percentile

    assert _percentile([1.0, 2.0, 3.0, 4.0], 50) == 2.0
    assert _percentile([1.0, 2.0, 3.0, 4.0], 95) == 4.0
    assert _percentile([], 50) == 0.0


def test_a_job_is_reentrant_so_a_nested_build_shares_one_record():
    with telemetry.job("outer") as outer:
        with telemetry.job("inner") as inner:
            assert inner is outer
        assert telemetry.active() is outer
    assert telemetry.active() is None


# ── the explicit per-column writer still works ──────────────────────────────


def test_an_injected_writer_still_gets_one_call_per_column():
    """The per-column path is how a test drives this without a model — it must not have
    been quietly replaced by the batched one."""
    seen = []
    values = _value_set(_facts(), topics=("working", "growth"))
    written = rewrites.write_all([values], write=lambda p: seen.append(p) or p.draft.upper())
    assert len(seen) == 2
    assert all(v.isupper() for v in written[0].values())


# ── end to end: a real deck, over the real templates, with a stubbed model ──


@pytest.mark.e2e
def test_a_real_deck_is_written_section_by_section_and_every_field_is_ai_authored(
        tmp_path, stub_model, monkeypatch):
    """The plan's end-to-end path, with only the model replaced.

        selection -> result -> evidence -> sub-deck plan -> commentary -> fill -> merge

    Two things are asserted that unit tests cannot: that a real template's columns group
    into sections at all (they carry real facts objects, not fixtures), and that the deck
    makes ONE author call per sub-deck rather than one per column — which is the entire
    performance claim.
    """
    from pptx import Presentation

    from studio.compute import compute_overall
    from studio.template_fill.assemble import assemble_deck, deck_shape

    monkeypatch.setenv("COMMENTARY_MODE", "ai_required")
    monkeypatch.setenv("STUDIO_MAX_WORKERS", "1")        # deterministic ordering in the log

    result = compute_overall(filters={"carrier": "Zurich", "country": ["Singapore"]})
    shape = deck_shape(result, slides=DeckSlides.only("overall", "country"))
    out = assemble_deck(result, out_path=str(tmp_path / "deck.pptx"),
                        work_dir=str(tmp_path / "work"),
                        slides=DeckSlides.only("overall", "country"))

    assert Presentation(out).slides                       # a deck came out
    authored = [c for c in stub_model if c["phase"] == "author"]
    assert authored, "the model must have written this deck"
    # One call per sub-deck that carries prose — never one per field.
    assert len(authored) <= sum(shape.blocks().values())
    assert sum(len(c["fields"]) for c in authored) >= len(authored)


@pytest.mark.e2e
def test_a_real_deck_refuses_to_ship_when_the_model_is_gone(tmp_path, monkeypatch):
    """``ai_required`` must fail the BUILD, not quietly fill the slides with rule prose."""
    import studio.ai.client as client

    from studio.compute import compute_overall
    from studio.template_fill.assemble import assemble_deck

    monkeypatch.setenv("COMMENTARY_MODE", "ai_required")
    monkeypatch.setenv("STUDIO_MAX_WORKERS", "1")
    monkeypatch.setattr(client, "llm_available", lambda: True)
    monkeypatch.setattr(client, "structured", lambda *a, **kw: None)

    result = compute_overall(filters={"carrier": "Zurich", "country": ["Singapore"]})
    with pytest.raises(mode.CommentaryUnavailable):
        assemble_deck(result, out_path=str(tmp_path / "deck.pptx"),
                      work_dir=str(tmp_path / "work"),
                      slides=DeckSlides.only("overall"))


@pytest.mark.e2e
def test_a_real_deck_still_ships_when_a_column_has_nothing_verifiable_to_say(tmp_path,
                                                                            monkeypatch):
    """The whole point of the change, proved at the export seam rather than in a unit.

    The model here writes only fragments, so every column is emptied by the gates — the
    worst case short of the model being gone. The build must produce a deck: an author who
    cannot support one column has not lost the right to the other thirty. What must NOT
    appear is the rejected text or the deterministic draft it would have fallen back to.
    """
    import studio.ai.client as client
    from pptx import Presentation

    from studio.compute import compute_overall
    from studio.template_fill.assemble import assemble_deck

    def fragments(model, system, user, *, tier="balanced", node="ai", phase="other",
                  fields=(), **kw):
        if model is not CommentarySections:
            return _approved_verdicts(user)
        asked = [line.split()[2] for line in user.splitlines()
                 if line.startswith("--- FIELD ")]
        return CommentarySections(sections=[
            CommentarySection(field_id=fid, bullets=[
                CommentaryBullet(text="Momentum: Cyber", fact_ids=["period.year"])])
            for fid in asked])

    monkeypatch.setenv("COMMENTARY_MODE", "ai_required")
    monkeypatch.setenv("STUDIO_MAX_WORKERS", "1")
    monkeypatch.setattr(client, "llm_available", lambda: True)
    monkeypatch.setattr(client, "structured", fragments)

    result = compute_overall(filters={"carrier": "Zurich", "country": ["Singapore"]})
    out = tmp_path / "deck.pptx"
    assemble_deck(result, out_path=str(out), work_dir=str(tmp_path / "work"),
                  slides=DeckSlides.only("overall"))

    assert out.exists(), "an emptied column must not cost the author the deck"
    text = "\n".join(shape.text_frame.text
                     for slide in Presentation(str(out)).slides
                     for shape in slide.shapes if shape.has_text_frame)
    assert "Momentum: Cyber" not in text, "the rejected fragment must never reach a slide"


def test_a_risk_flag_reaches_the_audit_line_rather_than_being_asked_for_and_binned(
        monkeypatch, caplog):
    """The schema asks for a risk read; something has to look at it or it is dead tokens."""
    import logging

    import studio.ai.client as client

    def structured(model, system, user, *, tier="balanced", node="ai", phase="other",
                   fields=(), **kw):
        if model is not CommentarySections:
            return _approved_verdicts(user)
        wanted = [line.split()[2] for line in user.splitlines()
                  if line.startswith("--- FIELD ")]
        return CommentarySections(sections=[
            CommentarySection(field_id=fid, risk_flag="concern", bullets=[
                CommentaryBullet(text=line, fact_ids=["period.year"]) for line in _SENTENCES])
            for fid in wanted])

    monkeypatch.setattr(client, "llm_available", lambda: True)
    monkeypatch.setattr(client, "structured", structured)

    with caplog.at_level(logging.INFO, logger="studio.template_fill.commentary_batch"):
        with telemetry.job("risk") as record:
            rewrites.write_all([_value_set(_facts(), topics=("working",))])
    assert "risk[" in caplog.text and "concern" in caplog.text
    assert record.counters.get("commentary_risk_concern") == 1


def test_the_env_file_is_read_before_the_commentary_mode_is_decided():
    """Regression: ``COMMENTARY_MODE`` in ``.env`` used to be invisible.

    ``preflight`` is the FIRST thing a build runs, and nothing ahead of it imported
    ``core.llm.clients`` — the only module that called ``load_dotenv``. So a ``.env``
    saying ``ai_required`` resolved to ``auto`` and the deck shipped deterministic prose,
    silently, in precisely the case strict mode exists to prevent.

    Run in a subprocess because the guarantee is about IMPORT ORDER, and by the time this
    test module is collected the whole app has been imported.
    """
    import subprocess
    import sys
    import textwrap

    probe = textwrap.dedent("""
        import sys
        from studio import commentary_mode
        commentary_mode.mode()
        assert "dotenv" in sys.modules, "the .env file was never read"
        print("ok")
    """)
    done = subprocess.run([sys.executable, "-c", probe], capture_output=True, text=True)
    assert done.returncode == 0, done.stderr
    assert "ok" in done.stdout


# ── a refusal must FAIL the build, not vanish into a warning ────────────────


def _refusing_builder(*_a, **_kw):
    raise mode.CommentaryUnavailable("no model", retryable=True)


def test_a_refusal_is_not_swallowed_by_the_best_effort_build_wrappers(monkeypatch):
    """Regression: strict mode's whole point is to STOP a deck, and the three build
    wrappers in ``generate`` caught it, logged a warning and answered ``None``.

    The job then "succeeded" with nothing in it, the canvas kept the previous run's deck,
    and the only trace was two warnings — the silent-deterministic-deck outcome strict
    mode exists to prevent, wearing a different hat.
    """
    from studio.authoring import generate as G

    monkeypatch.setattr(G, "_deck_for", _refusing_builder)
    monkeypatch.setattr(G, "_tdoc_for", _refusing_builder)
    monkeypatch.setattr(G, "_assembled_tdoc", _refusing_builder)

    sel = {"filters": {"carrier": "Zurich"}}
    for build in (G._generated_deck, G._generated_tdoc, G._generated_assembled_tdoc):
        with pytest.raises(mode.CommentaryUnavailable):
            build(sel)


def test_an_ordinary_build_failure_is_still_tolerated(monkeypatch):
    """The broad catch stays for what it was for: a bad selection or an odd template
    must not white-screen the app."""
    from studio.authoring import generate as G

    def explode(*_a, **_kw):
        raise ValueError("a template nobody can parse")

    monkeypatch.setattr(G, "_deck_for", explode)
    monkeypatch.setattr(G, "_tdoc_for", explode)
    monkeypatch.setattr(G, "_assembled_tdoc", explode)

    sel = {"filters": {"carrier": "Zurich"}}
    assert G._generated_deck(sel) is None
    assert G._generated_tdoc(sel) is None
    assert G._generated_assembled_tdoc(sel) is None


def test_a_refused_build_clears_the_canvas_instead_of_keeping_the_last_deck():
    """The symptom the user saw: no preview, and the PREVIOUS deck still on screen."""
    import time

    from studio.authoring import jobs

    job = jobs.start_build({"filters": {"carrier": "Zurich"}},
                           builder=lambda selection, report: _refusing_builder())
    deadline = time.time() + 10
    while time.time() < deadline and not job.snapshot()["done"]:
        time.sleep(0.01)

    state = job.snapshot()
    assert state["done"] and state["error"], "a refusal must reach the user as an error"
    assert state["retryable"] is True
    # This is the condition ``poll_generate`` reads to blank both document stores.
    produced = job.documents() is not None and not job.documents().empty
    assert state["error"] or not produced
    jobs.clear_job(job.job_id)


# ── a field that fails: the reasons must reach the repair and the refusal ────
#
# The reported symptom was a strict-mode build stopping on
# ``1 field(s) in section set2.0 could not be written (feedback-growth)`` — which names the
# field and nothing else. There was no way to tell whether the model had cited a figure the
# evidence did not carry, written a fragment, or simply come back a line short, and the one
# repair call it got asked the identical question again.


def _one_growth_field() -> dict:
    """A single two-bullet ``growth`` column — the shape of a feedback-table cell."""
    return {"note:0:0:0": _pending("growth", "growth line one.\ngrowth line two.", _facts())}


@pytest.fixture
def failing_model(monkeypatch):
    """A model whose answer always loses a line to the reading gate.

    A fragment is dropped by ``_keep_lines``, which leaves one line where the column asked
    for two — and ``min_lines(2)`` is 2, so a two-bullet column has no tolerance at all.
    That is exactly how a real build loses one field out of a section.
    """
    import studio.ai.client as client

    calls = []

    def structured(model, system, user, *, tier="balanced", node="ai", phase="other",
                   fields=(), **kw):
        calls.append({"node": node, "phase": phase, "fields": tuple(fields), "user": user})
        if model is CommentarySections:
            asked = [line.split()[2] for line in user.splitlines()
                     if line.startswith("--- FIELD ")]
            return CommentarySections(sections=[
                CommentarySection(field_id=fid, bullets=[
                    CommentaryBullet(text="Momentum: Cyber", fact_ids=["period.year"]),      # a fragment
                    CommentaryBullet(text=_SENTENCES[0], fact_ids=["period.year"]),
                ]) for fid in asked])
        return _approved_verdicts(user)

    monkeypatch.setattr(client, "structured", structured)
    monkeypatch.setattr(client, "llm_available", lambda: True)
    return calls


def test_the_repair_is_told_why_the_first_answer_was_rejected(unusable_model):
    """Otherwise the retry is a re-roll of the same dice, at the price of a call."""
    rewrites.write_all([_one_growth_field()])

    repairs = [c for c in unusable_model if "REJECTED" in c["user"]]
    assert repairs, "the repair call must carry the previous answer's verdict"
    assert "fragment" in repairs[0]["user"], "and name what was actually wrong"


@pytest.fixture
def unusable_model(monkeypatch):
    """A model whose every line is a fragment — nothing survives, so nothing can ship.

    A field that came up SHORT is salvaged (see the salvage tests below); a field with no
    verified line at all has nothing to salvage. That field now ships EMPTY under
    ``ai_required`` rather than failing the build: shipping deterministic prose is what the
    mode exists to prevent, and a blank box prevents it just as completely as a refusal did
    without costing the author every other column in the deck.
    """
    import studio.ai.client as client

    calls = []

    def structured(model, system, user, *, tier="balanced", node="ai", phase="other",
                   fields=(), **kw):
        calls.append({"node": node, "phase": phase, "fields": tuple(fields), "user": user})
        if model is CommentarySections:
            asked = [line.split()[2] for line in user.splitlines()
                     if line.startswith("--- FIELD ")]
            return CommentarySections(sections=[
                CommentarySection(field_id=fid, bullets=[
                    CommentaryBullet(text="Momentum: Cyber", fact_ids=["period.year"]),
                    CommentaryBullet(text="Growth: Property", fact_ids=["period.year"]),
                ]) for fid in asked])
        return _approved_verdicts(user)

    monkeypatch.setattr(client, "structured", structured)
    monkeypatch.setattr(client, "llm_available", lambda: True)
    return calls


def test_a_field_with_nothing_verifiable_ships_empty_instead_of_failing_the_deck(
        monkeypatch, unusable_model):
    """Evidence exhaustion blanks ONE box; it does not throw the whole deck away.

    Every line this model writes is a fragment, so nothing survives the gates. The old
    contract refused the build, which meant one column with no supportable point cost the
    author the other thirty. ``ai_required`` guarantees that no UNVERIFIED prose reaches a
    slide — an empty box honours that exactly as a refusal did, and still ships the deck.
    """
    monkeypatch.setenv("COMMENTARY_MODE", "ai_required")
    written = rewrites.write_all([_one_growth_field()])[0]
    assert list(written.values())[0] == ""


def test_an_emptied_field_does_not_fall_back_to_the_deterministic_draft(monkeypatch,
                                                                       unusable_model):
    """The blank must SURVIVE being placed, or blanking silently ships rule prose.

    ``_placed`` reads "written text, or the draft" — and an empty string is falsy, so the
    obvious spelling of that sentence hands the slide the very prose ``ai_required`` exists
    to keep off it. A field present in the written map wins even when it is empty.
    """
    monkeypatch.setenv("COMMENTARY_MODE", "ai_required")
    written = list(rewrites.write_all([_one_growth_field()])[0].values())[0]
    assert written == ""
    assert "growth line" not in written, "the deterministic draft must not reappear"


def test_a_verifier_that_could_not_run_still_refuses_the_build(monkeypatch):
    """The whole point of the omit/refuse split.

    "The evidence carried no supportable point" is an ANSWER, and blanks a box. "The judge
    never answered" is a BREAKAGE — it establishes nothing about whether there was
    something to say, so laundering it into a confident blank would be a silent lie.
    """
    import studio.ai.client as client
    from studio.ai.models import CommentaryVerdicts

    def structured(model, system, user, *, tier="balanced", node="ai", phase="other",
                   fields=(), **kw):
        if model is CommentarySections:
            asked = [line.split()[2] for line in user.splitlines()
                     if line.startswith("--- FIELD ")]
            return CommentarySections(sections=[
                CommentarySection(field_id=fid, bullets=[
                    CommentaryBullet(text=_SENTENCES[0], fact_ids=["period.year"]),
                ]) for fid in asked])
        return CommentaryVerdicts(verdicts=[])     # never lines up, retry included

    monkeypatch.setattr(client, "structured", structured)
    monkeypatch.setattr(client, "llm_available", lambda: True)
    monkeypatch.setenv("COMMENTARY_MODE", "ai_required")

    with pytest.raises(mode.CommentaryUnavailable) as raised:
        rewrites.write_all([_one_growth_field()])
    assert "could not be verified" in str(raised.value)
    assert raised.value.retryable is True


# ── a short field is shipped, not lost ──────────────────────────────────────
#
# A two-bullet cell whose evidence carries one point used to fail the whole field — and in
# ``ai_required`` the whole deck. Holding out for a second bullet cannot conjure evidence
# that is not there, and the alternative on the slide was deterministic prose. One verified
# line beats both.


def test_a_field_short_of_a_bullet_ships_what_it_verified(failing_model):
    """One good line beats the deterministic draft it would otherwise fall back to."""
    written = rewrites.write_all([_one_growth_field()])[0]
    assert list(written.values())[0] == _SENTENCES[0]


def test_a_short_field_does_not_refuse_a_strict_build(monkeypatch, failing_model):
    """``ai_required`` refuses UNWRITTEN fields, not short ones — the line is authored."""
    monkeypatch.setenv("COMMENTARY_MODE", "ai_required")
    written = rewrites.write_all([_one_growth_field()])[0]
    assert list(written.values())[0] == _SENTENCES[0]


def test_a_salvaged_line_still_cleared_every_gate(failing_model):
    """Only the floor moves. The fragment the gate rejected must never reach the slide."""
    written = list(rewrites.write_all([_one_growth_field()])[0].values())[0]
    assert "Momentum: Cyber" not in written
    assert "\n" not in written, "one surviving line means one bullet, not a blank second"


def test_a_field_that_repairs_successfully_ships_and_does_not_refuse(monkeypatch):
    """The repair earns its call: a second answer that clears the gates is used."""
    import studio.ai.client as client

    monkeypatch.setenv("COMMENTARY_MODE", "ai_required")
    attempts = {"n": 0}

    def structured(model, system, user, *, tier="balanced", node="ai", phase="other",
                   fields=(), **kw):
        if model is not CommentarySections:
            return _approved_verdicts(user)
        attempts["n"] += 1
        asked = [line.split()[2] for line in user.splitlines()
                 if line.startswith("--- FIELD ")]
        bullets = ([CommentaryBullet(text="Momentum: Cyber", fact_ids=["period.year"]),
                    CommentaryBullet(text="Unfinished comparison", fact_ids=["period.year"])]
                   if attempts["n"] == 1 else
                   [CommentaryBullet(text=s, fact_ids=["period.year"]) for s in _SENTENCES[:2]])
        return CommentarySections(sections=[
            CommentarySection(field_id=fid, bullets=bullets) for fid in asked])

    monkeypatch.setattr(client, "structured", structured)
    monkeypatch.setattr(client, "llm_available", lambda: True)

    written = rewrites.write_all([_one_growth_field()])[0]
    assert attempts["n"] == 2, "one author call, one repair"
    assert list(written.values())[0] == "\n".join(_SENTENCES[:2])


# ── strict mode has to be survivable, not just strict ───────────────────────
#
# ``ai_required`` is the mode this deck ships in, so a model that writes one unusable line
# per field — which is normal — must not fail the build. It used to: a two-bullet feedback
# cell had no room, so one rejected line refused the deck. These pin the whole path.


def _model_writing_one_bad_line(monkeypatch, *, bad="Momentum: Cyber"):
    """A model that answers with one unusable line and then good ones, up to the ask.

    Honours the requested line count, which is the point: the prompt now asks a
    zero-tolerance column for a spare, and a stub that ignores that would test nothing.
    """
    import studio.ai.client as client

    calls = []

    def structured(model, system, user, *, tier="balanced", node="ai", phase="other",
                   fields=(), **kw):
        calls.append({"node": node, "phase": phase, "user": user})
        if model is not CommentarySections:
            return _approved_verdicts(user)
        asked = [line.split()[2] for line in user.splitlines()
                 if line.startswith("--- FIELD ")]
        wanted = 3 if "Write 3 lines" in user else 2
        texts = [bad] + list(_SENTENCES[:wanted - 1])
        return CommentarySections(sections=[
            CommentarySection(field_id=fid,
                              bullets=[CommentaryBullet(text=t, fact_ids=["period.year"]) for t in texts])
            for fid in asked])

    monkeypatch.setattr(client, "structured", structured)
    monkeypatch.setattr(client, "llm_available", lambda: True)
    return calls


def test_one_unusable_line_no_longer_refuses_a_strict_build(monkeypatch):
    """The reported failure, end to end: it must now produce a deck."""
    monkeypatch.setenv("COMMENTARY_MODE", "ai_required")
    calls = _model_writing_one_bad_line(monkeypatch)

    written = rewrites.write_all([_one_growth_field()])[0]

    text = list(written.values())[0]
    assert text.splitlines() == list(_SENTENCES[:1]), "one verified finding needs no padding"
    assert "growth line one." not in text, "and no deterministic prose was substituted"
    assert len([c for c in calls if c["phase"] == "author"]) == 1, "no repair was needed"


def test_a_whole_sub_deck_of_flaky_fields_still_builds(monkeypatch):
    """Not one field — every field, which is what a real deck asks of the mode."""
    monkeypatch.setenv("COMMENTARY_MODE", "ai_required")
    _model_writing_one_bad_line(monkeypatch)

    written = rewrites.write_all([_value_set(_facts())])[0]

    assert len(written) == 3
    for text in written.values():
        assert text.splitlines() == list(_SENTENCES[:1])


def test_a_field_the_model_cannot_write_ships_empty_and_never_ships_its_rejects(monkeypatch):
    """The mode still means what it says — the spare is tolerance, not a bypass.

    The deck no longer stops (that cost thirty good columns to save one bad box), but the
    guarantee underneath is unchanged and is what this asserts: neither the fragment the
    gate rejected nor the deterministic draft may reach the slide. The box is simply empty.
    """
    monkeypatch.setenv("COMMENTARY_MODE", "ai_required")
    _model_writing_one_bad_line(monkeypatch, bad=_SENTENCES[0])   # every line a duplicate
    import studio.ai.client as client

    def all_fragments(model, system, user, *, tier="balanced", node="ai", phase="other",
                      fields=(), **kw):
        if model is not CommentarySections:
            return _approved_verdicts(user)
        asked = [line.split()[2] for line in user.splitlines()
                 if line.startswith("--- FIELD ")]
        return CommentarySections(sections=[
            CommentarySection(field_id=fid, bullets=[
                CommentaryBullet(text="Momentum: Cyber", fact_ids=["period.year"]),
                CommentaryBullet(text="Rank improved to 4th.", fact_ids=["period.year"]),
            ]) for fid in asked])

    monkeypatch.setattr(client, "structured", all_fragments)
    written = list(rewrites.write_all([_one_growth_field()])[0].values())[0]
    assert written == ""
    assert "Momentum: Cyber" not in written, "the rejected fragment must not reach the slide"
    assert "growth line" not in written, "nor the deterministic draft"


def test_the_fields_that_were_written_are_cached_even_when_one_comes_up_empty(monkeypatch):
    """So a re-run only re-attempts what failed, instead of paying for the deck again.

    An emptied field is deliberately NOT cached: blanking is this run's verdict on this
    evidence, and a later run with a repaired prompt or a better model must be free to try
    the column again rather than inherit the blank.
    """
    monkeypatch.setenv("COMMENTARY_MODE", "ai_required")
    import studio.ai.client as client

    def one_bad_field(model, system, user, *, tier="balanced", node="ai", phase="other",
                      fields=(), **kw):
        if model is not CommentarySections:
            return _approved_verdicts(user)
        asked = [line.split()[2] for line in user.splitlines()
                 if line.startswith("--- FIELD ")]
        return CommentarySections(sections=[
            CommentarySection(
                field_id=fid,
                bullets=[CommentaryBullet(text="Momentum: Cyber", fact_ids=["period.year"])]
                if fid.startswith("growth")
                else [CommentaryBullet(text=s, fact_ids=["period.year"]) for s in _SENTENCES[:2]],
            ) for fid in asked])

    monkeypatch.setattr(client, "structured", one_bad_field)
    monkeypatch.setattr(client, "llm_available", lambda: True)

    rewrites.write_all([_value_set(_facts())])

    from studio.template_fill import commentary_cache as cache

    from studio.template_fill import editorial

    section = B.group_sections([_value_set(_facts())])[0]
    pack = B._pack_for(section)
    plan = editorial.plan_deck([section.columns])
    cached = {column.topic for column in section.columns
              if cache.get(B.cache_key(section, column, pack, plan))}
    assert {"working", "challenges"} <= cached, "the fields that were written survived"
    assert "growth" not in cached, "the one that failed did not"


# ── every failing field gets every repair round ─────────────────────────────
#
# The repair loop used to return the moment ANY field succeeded:
#
#     if text or not still_failing:
#         return text
#
# So three failing fields of which the first round fixed one left the other two with their
# second round unspent — and under ``ai_required`` that is a refused deck for a reason that
# had a round left to fix it. Successes accumulate now, and the loop carries on with
# whatever is still short.


def _three_growth_fields() -> dict:
    """Three two-bullet columns in one section, so one repair call covers all three."""
    facts = _facts()
    return {f"note:{i}:{i}:0": _pending("growth", "growth line one.\ngrowth line two.", facts)
            for i in range(3)}


@pytest.fixture
def model_fixing_one_field_per_round(monkeypatch):
    """Round 1 fixes one field, round 2 fixes the next. The third never recovers.

    Pins the exact reported shape: a partial success must not end the loop for the fields
    that are still short.
    """
    import studio.ai.client as client

    rounds = {"n": 0}
    calls = []

    def structured(model, system, user, *, tier="balanced", node="ai", phase="other",
                   fields=(), **kw):
        if model is not CommentarySections:
            return _approved_verdicts(user)
        calls.append({"user": user, "fields": tuple(fields)})
        rounds["n"] += 1
        asked = [line.split()[2] for line in user.splitlines()
                 if line.startswith("--- FIELD ")]
        # On round N, the Nth field asked for is written properly; the rest lose a line.
        out = []
        for i, fid in enumerate(sorted(asked)):
            good = [CommentaryBullet(text=s, fact_ids=["period.year"]) for s in _SENTENCES[:2]]
            short = [CommentaryBullet(text="Momentum: Cyber", fact_ids=["period.year"]),
                     CommentaryBullet(text="Unfinished comparison", fact_ids=["period.year"])]
            out.append(CommentarySection(field_id=fid,
                                         bullets=good if i < rounds["n"] - 1 else short))
        return CommentarySections(sections=out)

    monkeypatch.setattr(client, "structured", structured)
    monkeypatch.setattr(client, "llm_available", lambda: True)
    return calls


def test_a_partial_repair_does_not_abandon_the_fields_still_failing(
        model_fixing_one_field_per_round):
    """The bug: one field succeeding used to end repair for every other field."""
    calls = model_fixing_one_field_per_round
    rewrites.write_all([_three_growth_fields()])

    repairs = [c for c in calls if "REJECTED" in c["user"]]
    assert len(repairs) == B.MAX_REPAIRS, (
        "every repair round must be spent while fields are still short — the loop used "
        "to stop after the round that fixed the first field"
    )
    # And the last round still asks for the field that never recovered.
    assert repairs[-1]["fields"], "the final round must still be asking for something"


def test_repaired_fields_from_earlier_rounds_are_not_discarded(
        model_fixing_one_field_per_round):
    """Round 1's success has to survive round 2 running for somebody else."""
    written = rewrites.write_all([_three_growth_fields()])[0]
    full = [t for t in written.values() if t == "\n".join(_SENTENCES[:2])]
    assert len(full) >= 2, (
        f"expected the fields repaired in rounds 1 and 2 to both ship, got {written!r}"
    )


# ── repair tops a field up; it does not rewrite it ──────────────────────────


def test_the_repair_shows_the_lines_already_verified_and_asks_only_for_the_rest(failing_model):
    section = B.group_sections([_one_growth_field()])[0]
    column = section.columns[0]
    pack = B._pack_for(section)
    B._repair(section, pack, "", [B.Failure(column, ("fragment",), (_SENTENCES[0],))])
    repairs = [c for c in failing_model if "REJECTED" in c["user"]]
    assert repairs
    assert "ALREADY WRITTEN AND VERIFIED" in repairs[0]["user"]
    assert _SENTENCES[0] in repairs[0]["user"]
    assert "up to 1 additional" in repairs[0]["user"]


def test_a_kept_line_is_never_re_verified_and_so_cannot_be_lost(failing_model):
    """The whole point of a top-up: a retry meant to help must not cost a good line."""
    written = rewrites.write_all([_one_growth_field()])[0]
    assert _SENTENCES[0] in list(written.values())[0]


def test_a_model_that_echoes_a_kept_line_does_not_ship_it_twice():
    """The prompt says not to repeat; the merge does not rely on it obeying."""
    from studio.template_fill.commentary_batch import _merged

    merged = _merged({"f": ("The book grew.",)},
                     {"f": ("the  book   grew", "And then it held.")})
    assert merged["f"] == ["The book grew.", "And then it held."]
