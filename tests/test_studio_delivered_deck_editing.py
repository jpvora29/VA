"""Editing the delivered deck's own words: click the slide, type in the edit field.

An assembled QBR is a finished ``.pptx`` with no slots left to fill, so a text box
is addressed by where it sits — ``slide:shape`` — and the whole block of lines is
what an author edits. These tests walk the path end to end: the address rules, the
click target the canvas puts over each text box, the edit field, what the canvas
reflects back, and the PowerPoint the author downloads.
"""
from __future__ import annotations

import pathlib

import pytest
from dash.development.base_component import Component
from pptx import Presentation

from studio.page import template_preview as TP
from studio.template_fill import text_edits as TE
from studio.template_fill.analyze import analyze
from studio.template_fill.fill import apply_text_overrides

TEMPLATE = pathlib.Path(__file__).parents[1] / "template" / "overall_template.pptx"


@pytest.fixture(scope="module")
def template():
    if not TEMPLATE.exists():
        pytest.skip("the overall template is not checked out")
    return analyze(str(TEMPLATE))


def _tdoc(**extra):
    """A delivered (assembled) document — filled, with a diagnostic-only manifest."""
    return {
        "template_path": str(TEMPLATE),
        "values": {}, "manifest": [], "overrides": {}, "map_overrides": {},
        "added": {}, "hidden": [], "order": list(range(6)),
        "width_emu": 12192000, "height_emu": 6858000, "n_slides": 6,
        "assembled": True, "background_urls": [None] * 6,
        **extra,
    }


def _rendered_tdoc(**extra):
    """The same document as the author sees it: every slide an exact PNG render."""
    return _tdoc(background_urls=[f"/assets/render-{i}.png" for i in range(6)], **extra)


def _walk(node):
    if node is None:
        return
    if isinstance(node, (list, tuple)):
        for child in node:
            yield from _walk(child)
        return
    yield node
    if isinstance(node, Component):
        yield from _walk(getattr(node, "children", None))


def _of_type(node, kind):
    return [
        item for item in _walk(node)
        if isinstance(getattr(item, "id", None), dict) and item.id.get("type") == kind
    ]


def _classes(node):
    out = set()
    for item in _walk(node):
        if getattr(item, "className", None):
            out.update(str(item.className).split())
    return out


def _prose_box(template):
    """The first shape on the deck holding several paragraphs of prose."""
    for slide in template.slides:
        for shape in slide.shapes:
            block = TE.editable_text(shape, slide.index, {})
            if block and len(block.lines) >= 2 and all(len(x) > 40 for x in block.lines[:2]):
                return slide, shape, block
    pytest.skip("no multi-paragraph prose box in this template")


def _leaves(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _leaves(shape.shapes)
        else:
            yield shape


# ── the address ──────────────────────────────────────────────────────────────


def test_a_text_box_is_addressed_by_where_it_sits():
    assert TE.parse_address("3:10") == TE.ShapeAddress(3, 10)
    assert TE.ShapeAddress(3, 10).key == "3:10"


@pytest.mark.parametrize("key", ["", None, "3", "3:10:0", "a:10", "3:-1", "note:3:10"])
def test_a_malformed_address_is_refused(key):
    assert TE.parse_address(key) is None


# ── what is offered for editing ──────────────────────────────────────────────


def test_a_prose_box_offers_all_of_its_lines(template):
    _slide, _shape, block = _prose_box(template)

    assert len(block.lines) >= 2
    assert block.lines == block.original
    assert block.edited is False
    assert block.text == "\n".join(block.lines)


def test_data_shapes_and_blank_spacing_are_not_prose(template):
    kinds = {shape.kind for slide in template.slides for shape in slide.shapes}
    assert kinds - {"text"}, "this template should have non-text shapes to exclude"

    for slide in template.slides:
        for shape in slide.shapes:
            block = TE.editable_text(shape, slide.index, {})
            if shape.kind != "text":
                assert block is None, f"{shape.kind} is not prose"
            if block:
                assert all(line.strip() for line in block.lines)


def test_an_edited_box_reports_itself_and_shows_the_new_words(template):
    slide, shape, block = _prose_box(template)

    edited = TE.editable_text(
        shape, slide.index, {block.address.key: ["Marine carried the quarter."]}
    )

    assert edited.lines == ("Marine carried the quarter.",)
    assert edited.original == block.original
    assert edited.edited is True


# ── committing an edit ───────────────────────────────────────────────────────


def test_the_edit_field_splits_what_was_typed_into_paragraphs():
    doc = TE.set_text_edit(
        _tdoc(), "3:10", "  First line. \n\n  Second line.  \n", original=["Old"]
    )

    assert TE.text_edits(doc) == {"3:10": ["First line.", "Second line."]}


def test_adding_and_deleting_lines_is_just_editing_the_block():
    doc = TE.set_text_edit(_tdoc(), "3:10", "A\nB\nC", original=["A", "B"])
    assert TE.text_edits(doc)["3:10"] == ["A", "B", "C"]

    doc = TE.set_text_edit(doc, "3:10", "A", original=["A", "B"])
    assert TE.text_edits(doc)["3:10"] == ["A"]


def test_typing_the_decks_own_words_back_clears_the_edit():
    doc = TE.set_text_edit(_tdoc(), "3:10", "New", original=["Old"])
    doc = TE.set_text_edit(doc, "3:10", "Old", original=["Old"])

    assert TE.text_edits(doc) == {}


def test_emptying_the_field_is_a_reset_not_a_blank_box():
    """A text box emptied to nothing would delete the deck's prose; it reverts."""
    doc = TE.set_text_edit(_tdoc(), "3:10", "New", original=["Old"])
    doc = TE.set_text_edit(doc, "3:10", "   \n  ", original=["Old"])

    assert TE.text_edits(doc) == {}


def test_reset_puts_a_box_back_to_what_the_deck_says():
    doc = TE.set_text_edit(_tdoc(), "3:10", "New", original=["Old"])

    assert TE.text_edits(TE.clear_text_edit(doc, "3:10")) == {}


def test_an_unaddressable_key_cannot_write_to_the_document():
    doc = TE.set_text_edit(_tdoc(), "template_path", "/etc/passwd", original=[])

    assert TE.text_edits(doc) == {}
    assert doc["template_path"] == str(TEMPLATE)


def test_edits_group_the_way_the_writer_walks_the_deck():
    grouped = TE.grouped({"3:10": ["a", "b"], "4:7": ["c"], "junk": ["d"]})

    assert grouped == {3: {10: ["a", "b"]}, 4: {7: ["c"]}}


# ── the canvas: a click target, and only then a reflection ───────────────────


def test_every_text_box_on_the_rendered_slide_can_be_clicked(template):
    slide, _shape, block = _prose_box(template)
    tdoc = _rendered_tdoc()

    body = TP.template_preview_body(tdoc, {"idx": tdoc["order"].index(slide.index)})
    picks = {item.id["at"] for item in _of_type(body, "qs-tf-pick")}

    assert block.address.key in picks
    assert len(picks) > 1, "every text box on the page is a target, not just one"


def test_an_unedited_slide_draws_nothing_over_the_render(template):
    """Drawing the words a second time is what made the slide look broken."""
    slide, _shape, _block = _prose_box(template)
    tdoc = _rendered_tdoc()

    body = TP.template_preview_body(tdoc, {"idx": tdoc["order"].index(slide.index)})

    assert "qs-tf-reflect" not in _classes(body)
    assert "is-edited" not in _classes(body)


def test_the_canvas_reflects_what_was_typed(template):
    slide, _shape, block = _prose_box(template)
    tdoc = TE.set_text_edit(
        _rendered_tdoc(), block.address.key,
        "Marine carried the quarter.\nProperty needs a plan.", original=block.original,
    )

    body = TP.template_preview_body(tdoc, {"idx": tdoc["order"].index(slide.index)})
    shown = [
        item.children for item in _walk(body)
        if "qs-tf-reflect-line" in str(getattr(item, "className", ""))
    ]

    assert shown == ["Marine carried the quarter.", "Property needs a plan."]
    assert "is-edited" in _classes(body)
    pills = [
        str(getattr(item, "children", ""))
        for item in _walk(body)
        if "qs-tf-pill" in str(getattr(item, "className", ""))
    ]
    assert any("1 edited" in pill for pill in pills), "the slide bar counts the edits"


def test_clicking_a_box_opens_every_line_in_its_own_field(template):
    """One run-on textarea is what made the field unusable; a line is a line."""
    slide, _shape, block = _prose_box(template)
    tdoc = _rendered_tdoc()
    view = {"idx": tdoc["order"].index(slide.index), "tf_sel": block.address.key}

    body = TP.template_preview_body(tdoc, view)
    fields = _of_type(body, "qs-tf-line")

    assert [f.value for f in fields] == list(block.lines)
    assert [f.id["i"] for f in fields] == list(range(len(block.lines)))
    assert all(f.id["at"] == block.address.key for f in fields)
    assert len(_of_type(body, "qs-tf-linedel")) == len(block.lines), "each line can go"
    assert len(_of_type(body, "qs-tf-lineadd")) == 1


def test_the_edit_field_waits_until_something_is_selected(template):
    slide, _shape, _block = _prose_box(template)
    tdoc = _rendered_tdoc()

    body = TP.template_preview_body(tdoc, {"idx": tdoc["order"].index(slide.index)})

    assert _of_type(body, "qs-tf-line") == []
    assert any(
        "Click any text on the slide" in str(getattr(item, "children", ""))
        for item in _walk(body)
    )


def test_a_selection_from_another_page_does_not_follow_you(template):
    slide, _shape, block = _prose_box(template)
    other = next(s for s in template.slides if s.index != slide.index)
    tdoc = _rendered_tdoc()
    view = {"idx": tdoc["order"].index(other.index), "tf_sel": block.address.key}

    assert _of_type(TP.template_preview_body(tdoc, view), "qs-tf-line") == []


def test_the_reset_button_is_live_only_on_a_box_that_was_edited(template):
    slide, _shape, block = _prose_box(template)
    tdoc = _rendered_tdoc()
    view = {"idx": tdoc["order"].index(slide.index), "tf_sel": block.address.key}

    clean = _of_type(TP.template_preview_body(tdoc, view), "qs-tf-reset")[0]
    edited_doc = TE.set_text_edit(tdoc, block.address.key, "New words.", original=block.original)
    dirty = _of_type(TP.template_preview_body(edited_doc, view), "qs-tf-reset")[0]

    assert clean.disabled is True
    assert dirty.disabled is False


def test_no_editor_field_declares_a_prop_dash_will_reject(template):
    """`dcc.Textarea.spellCheck` is a boolean; a string there is a runtime error."""
    from dash import dcc

    slide, _shape, block = _prose_box(template)
    tdoc = _rendered_tdoc()
    view = {"idx": tdoc["order"].index(slide.index), "tf_sel": block.address.key}

    for field in _of_type(TP.template_preview_body(tdoc, view), "qs-tf-line"):
        assert isinstance(field, dcc.Textarea)
        assert not isinstance(getattr(field, "spellCheck", None), str)


# ── the edit field repaints on its own ───────────────────────────────────────


def test_the_edit_field_can_be_built_without_rebuilding_the_slide(template):
    """Selection is repainted alone — going through the master render put the
    "Opening…" overlay up for what is only a click."""
    slide, _shape, block = _prose_box(template)
    tdoc = _rendered_tdoc()
    view = {"idx": tdoc["order"].index(slide.index)}

    panel = TP.text_editor_for(tdoc, view, block.address.key)
    fields = _of_type(panel, "qs-tf-line")

    assert [f.value for f in fields] == list(block.lines)


def test_the_edit_field_survives_a_document_it_cannot_read():
    assert _of_type(TP.text_editor_for(None, {}, "3:10"), "qs-tf-line") == []
    assert _of_type(TP.text_editor_for({"template_path": "/nope.pptx"}, {}, "3:10"),
                    "qs-tf-line") == []


# ── line by line ─────────────────────────────────────────────────────────────


def test_retyping_one_line_leaves_its_neighbours_alone():
    doc = TE.set_line(_tdoc(), "3:10", 1, "Rewritten.", original=["A", "B", "C"])

    assert TE.text_edits(doc)["3:10"] == ["A", "Rewritten.", "C"]


def test_adding_a_line_opens_a_blank_one_to_write_in():
    doc = TE.add_line(_tdoc(), "3:10", original=["A", "B"])

    assert TE.text_edits(doc)["3:10"] == ["A", "B", ""]
    assert TE.grouped(TE.text_edits(doc)) == {3: {10: ["A", "B"]}}, "a blank line is not a paragraph"


def test_adding_twice_without_writing_does_not_stack_blank_lines():
    doc = TE.add_line(_tdoc(), "3:10", original=["A"])

    assert TE.text_edits(TE.add_line(doc, "3:10", original=["A"]))["3:10"] == ["A", ""]


def test_deleting_a_line_removes_that_one():
    doc = TE.delete_line(_tdoc(), "3:10", 0, original=["A", "B", "C"])

    assert TE.text_edits(doc)["3:10"] == ["B", "C"]


def test_deleting_back_to_the_decks_own_words_clears_the_edit():
    doc = TE.add_line(_tdoc(), "3:10", original=["A", "B"])
    doc = TE.delete_line(doc, "3:10", 2, original=["A", "B"])

    assert TE.text_edits(doc) == {}


def test_deleting_every_line_is_a_reset_not_a_blank_box():
    """Emptying a box entirely would delete the deck's prose; it reverts instead."""
    doc = _tdoc()
    for _ in range(3):
        doc = TE.delete_line(doc, "3:10", 0, original=["A", "B", "C"])

    assert TE.text_edits(doc) == {}


@pytest.mark.parametrize("index", [-1, 9])
def test_a_line_that_is_not_there_changes_nothing(index):
    original = ["A", "B"]

    assert TE.text_edits(TE.set_line(_tdoc(), "3:10", index, "x", original=original)) == {}
    assert TE.text_edits(TE.delete_line(_tdoc(), "3:10", index, original=original)) == {}


def test_an_unaddressable_key_cannot_write_a_line():
    doc = TE.set_line(_tdoc(), "template_path", 0, "/etc/passwd", original=["A"])

    assert TE.text_edits(doc) == {}
    assert doc["template_path"] == str(TEMPLATE)


# ── the colours a reflected box is painted in ────────────────────────────────


def test_a_reflected_box_is_painted_in_the_slide_s_own_colours(template):
    """It covers the render, so it must not arrive as a white box on a navy page."""
    dark = next(s for s in template.slides if s.background_color == "000F47")
    light = next(s for s in template.slides if s.background_color == "CEECFF")

    dark_title = next(s for s in dark.shapes if s.kind == "text" and s.text.strip())
    light_text = next(s for s in light.shapes if s.kind == "text" and s.text.strip())

    assert TP._ink(dark_title, dark) == ("#000F47", "#ffffff")
    assert TP._ink(light_text, light) == ("#CEECFF", "#0b1f44")


def test_the_paper_is_the_nearest_backdrop_the_text_sits_on(template):
    """A divider headline has no fill of its own — the full-bleed panel behind it does."""
    divider = next(s for s in template.slides
                   if s.background_color == "000F47" and any(sh.fill_color for sh in s.shapes))
    panel = next(sh for sh in divider.shapes if sh.fill_color)
    over_it = next(sh for sh in divider.shapes
                   if sh is not panel and sh.kind == "text" and not sh.fill_color
                   and TP._contains(panel, sh))

    assert TP._panel_fill(over_it, divider) == panel.fill_color


def test_a_slide_inherits_the_paper_its_layout_sets(template):
    """Nothing may report a blank background just because the fill lives on the layout."""
    assert all(slide.background_color for slide in template.slides)


# ── end to end: canvas click → edit field → the downloaded PowerPoint ────────


def test_what_was_typed_reaches_the_downloaded_deck(tmp_path, template):
    slide, shape, block = _prose_box(template)
    tdoc = TE.set_text_edit(
        _rendered_tdoc(), block.address.key,
        "Marine cross-sell carried the quarter.\nProperty retention needs a plan.",
        original=block.original,
    )

    out = tmp_path / "delivered-edited.pptx"
    apply_text_overrides(str(TEMPLATE), TE.text_edits(tdoc), str(out))
    box = next(s for s in _leaves(Presentation(out).slides[slide.index].shapes)
               if int(s.shape_id) == shape.shape_id)

    assert [p.text for p in box.text_frame.paragraphs] == [
        "Marine cross-sell carried the quarter.",
        "Property retention needs a plan.",
    ]
    # This template repeats its example prose in a second commentary column; that
    # copy must be untouched, which is what proves an address names ONE box.
    others = [s for s in _leaves(Presentation(out).slides[slide.index].shapes)
              if hasattr(s, "text") and int(s.shape_id) != shape.shape_id]
    assert any(block.original[0] in s.text for s in others)


def test_a_line_added_in_the_edit_field_reaches_the_deck(tmp_path, template):
    slide, shape, block = _prose_box(template)
    added = list(block.original) + ["And one the author added."]
    tdoc = TE.set_text_edit(_rendered_tdoc(), block.address.key,
                            "\n".join(added), original=block.original)

    out = tmp_path / "added.pptx"
    apply_text_overrides(str(TEMPLATE), TE.text_edits(tdoc), str(out))
    box = next(s for s in _leaves(Presentation(out).slides[slide.index].shapes)
               if int(s.shape_id) == shape.shape_id)

    assert [p.text for p in box.text_frame.paragraphs] == added


def test_a_line_deleted_in_the_edit_field_leaves_the_deck(tmp_path, template):
    slide, shape, block = _prose_box(template)
    kept = list(block.original[:1])
    tdoc = TE.set_text_edit(_rendered_tdoc(), block.address.key,
                            "\n".join(kept), original=block.original)

    out = tmp_path / "deleted.pptx"
    apply_text_overrides(str(TEMPLATE), TE.text_edits(tdoc), str(out))
    box = next(s for s in _leaves(Presentation(out).slides[slide.index].shapes)
               if int(s.shape_id) == shape.shape_id)

    assert [p.text for p in box.text_frame.paragraphs] == kept


def test_an_edit_keeps_the_box_s_own_formatting(tmp_path, template):
    """The run-preserving write is what keeps think-cell and run styling alive."""
    slide, shape, block = _prose_box(template)

    def font_of(path, para=0):
        box = next(s for s in _leaves(Presentation(path).slides[slide.index].shapes)
                   if int(s.shape_id) == shape.shape_id)
        run = box.text_frame.paragraphs[para].runs[0]
        return (run.font.name, run.font.size, run.font.bold,
                getattr(run.font.color, "rgb", None))

    out = tmp_path / "formatted.pptx"
    apply_text_overrides(
        str(TEMPLATE),
        {block.address.key: ["Rewritten line.", "Second line.", "A third the author added."]},
        str(out),
    )

    assert font_of(out) == font_of(TEMPLATE)
    assert font_of(out, 2) == font_of(TEMPLATE), "an added line looks like its neighbours"


def test_the_build_on_disk_is_never_rewritten(tmp_path, template):
    _slide, _shape, block = _prose_box(template)
    before = TEMPLATE.read_bytes()

    apply_text_overrides(str(TEMPLATE), {block.address.key: ["Rewritten."]},
                         str(tmp_path / "copy.pptx"))

    assert TEMPLATE.read_bytes() == before


def test_an_address_that_no_longer_exists_is_skipped_quietly(tmp_path):
    out = tmp_path / "ghost.pptx"

    apply_text_overrides(str(TEMPLATE), {"99:99": ["nowhere"]}, str(out))

    assert out.exists(), "a stale edit must not cost the author the export"


# ── the export the author clicks ─────────────────────────────────────────────


def test_export_serves_the_cached_build_when_nothing_was_edited():
    """An untouched deck must not pay for a rewrite, nor get a different file."""
    from studio.authoring.export import _with_retyped_text

    assert _with_retyped_text(_tdoc()) == str(TEMPLATE)


def test_export_serves_a_copy_carrying_the_edits(template):
    from studio.authoring.export import _delivered_text, _with_retyped_text

    _slide, shape, block = _prose_box(template)
    tdoc = _tdoc()
    original = _delivered_text(tdoc, block.address.key)
    assert original == block.original, "the export reads the same words the canvas shows"

    tdoc = TE.set_text_edit(tdoc, block.address.key, "Rewritten for the board.",
                            original=original)
    out = _with_retyped_text(tdoc)

    assert out != str(TEMPLATE), "the build on disk is never the edited file"
    box = next(s for s in _leaves(Presentation(out).slides[block.address.slide_idx].shapes)
               if int(s.shape_id) == shape.shape_id)
    assert [p.text for p in box.text_frame.paragraphs] == ["Rewritten for the board."]


def test_resetting_every_edit_sends_the_build_again(template):
    from studio.authoring.export import _with_retyped_text

    _slide, _shape, block = _prose_box(template)
    tdoc = TE.set_text_edit(_tdoc(), block.address.key, "New.", original=block.original)

    assert _with_retyped_text(TE.clear_text_edit(tdoc, block.address.key)) == str(TEMPLATE)
