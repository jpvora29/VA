"""The explainable Boardroom, end to end.

The business flow this covers is the one the roadmap's acceptance criteria
describe:

    digest -> editable document -> rendered board -> edited widget -> PPTX

with the invariants that make the board defensible checked at each hop: no
widget shows a score or a severity, every comparison names its basis, a widget
the data cannot support is dropped rather than explained, no panel appears
twice, and the exported slide states the same scope and the same values as the
screen.

Run:  pytest tests/test_boardroom_explainable.py -q -o pythonpath=.
"""
from __future__ import annotations

import copy

import pytest
from dash.development.base_component import Component

from core.boardroom import derive, opportunity
from ui.boardroom import builder, catalog, editor, model, ppt_export, widgets_generated
from ui.boardroom.render import render_document

EXPLAINABLE_KINDS = (
    "watchlist",
    "headroom",
    "whitespace",
    "quarterly",
    "portfolio_map",
    "top_carriers",
)


# ── a realistic digest, in the shape `boardroom_node` emits ──────────────────


def digest() -> dict:
    return {
        "title": "Zurich — Canada",
        "subtitle": "Q2 2026 premium performance",
        "headline": "Property contraction cost £3.2m in the quarter.",
        "kpis": [{"label": "Premium", "value": "£19.2m", "delta": "-14.3% QoQ", "tone": "danger"}],
        "commentary": [{"heading": "The big picture", "points": ["Premium fell in Property."]}],
        "risks": [{"label": "Rank decline", "severity": "High", "tone": "danger"}],
        "scope": [
            {"key": "country", "label": "Country", "value": "Canada", "icon": "bi bi-globe2",
             "source": "asked in this question"},
            {"key": "product", "label": "Product", "value": "Property", "icon": "bi bi-box-seam",
             "source": "asked in this question"},
        ],
        "watchlist": {
            "items": [
                {
                    "risk": "Property premium contraction",
                    "scope": "Canada / Property / Zurich",
                    "premium_exposed": "£12.4m",
                    "premium_exposed_value": 12_400_000,
                    "movement": "-14.2% QoQ",
                    "movement_pct": -14.2,
                    "adverse": True,
                    "comparison": "Q2 2026 vs Q1 2026",
                    "trigger": "Premium fell 14.2% between Q1 and Q2 2026",
                    "consecutive_periods": 2,
                    "periods_comparable": True,
                    "owner_action": "Review top-lost industries with Placement",
                    "tone": "danger",
                }
            ],
            "basis": "Q2 2026 vs Q1 2026",
        },
        "headroom": {
            "rows": [
                {
                    "product_line": "Property",
                    "carrier_premium": "£8.2m", "carrier_premium_value": 8_200_000,
                    "marsh_premium": "£42.0m", "marsh_premium_value": 42_000_000,
                    "share_of_wallet_pct": 19.5, "share_of_portfolio_pct": 46.0,
                    "whitespace_premium": "£33.8m", "whitespace_premium_value": 33_800_000,
                    "market_change": "+6.4% QoQ", "status": "established",
                    "focus": "Expand in Manufacturing",
                }
            ],
            "definition": "Whitespace premium = Marsh premium - carrier premium (floored at zero).",
            "basis": "Q2 2026",
        },
        "whitespace": {
            "rows": [
                {
                    "industry": "Manufacturing", "product_line": "Property",
                    "marsh_premium": "£18.6m", "marsh_premium_value": 18_600_000,
                    "carrier_premium": "£0.0m", "carrier_premium_value": 0.0,
                    "share_of_wallet_pct": 0.0, "share_of_portfolio_pct": 0.0,
                    "peer_share_of_wallet_pct": 34.0,
                    "whitespace_premium": "£18.6m", "whitespace_premium_value": 18_600_000,
                    "marsh_change": "+8.2% QoQ", "status": "no_premium",
                    "focus_reason": "Largest unwritten premium and the fastest-growing industry",
                },
                {
                    "industry": "Construction", "product_line": "Cyber",
                    "marsh_premium": "£15.2m", "marsh_premium_value": 15_200_000,
                    "carrier_premium": "£0.9m", "carrier_premium_value": 900_000,
                    "share_of_wallet_pct": 5.9,
                    "whitespace_premium": "£14.3m", "whitespace_premium_value": 14_300_000,
                    "marsh_change": "+3.1% QoQ", "status": "low_presence",
                },
            ],
            "product_line": "Property",
            "product_lines": ["Property", "Cyber"],
            "definition": "Marsh premium above materiality, carrier premium zero or low.",
        },
        "quarterly": {
            "rows": [
                {"quarter": "Q1 2026", "premium": "£22.4m", "premium_value": 22_400_000,
                 "change_currency": "+£2.3m", "change_pct": 11.4, "share_of_wallet_pct": 13.5,
                 "driver": "Property growth", "complete": True},
                {"quarter": "Q2 2026", "premium": "£19.2m", "premium_value": 19_200_000,
                 "change_currency": "-£3.2m", "change_pct": -14.3, "share_of_wallet_pct": 11.9,
                 "driver": "Construction decline", "complete": True},
            ],
            "basis": "QoQ",
        },
        "portfolio_map": {
            "bubbles": [
                {"product_line": "Property", "share_of_wallet_pct": 19.5,
                 "share_of_portfolio_pct": 46.0, "premium": "£8.2m", "premium_value": 8_200_000,
                 "marsh_premium": "£42.0m", "marsh_premium_value": 42_000_000,
                 "growth": "+6.4% YoY", "growth_pct": 6.4, "peer_share_of_wallet_pct": 22.0,
                 "tone": "neutral"},
                {"product_line": "Cyber", "share_of_wallet_pct": 3.4,
                 "share_of_portfolio_pct": 2.2, "premium": "£0.4m", "premium_value": 400_000,
                 "marsh_premium": "£11.7m", "marsh_premium_value": 11_700_000,
                 "growth": "+18.1% YoY", "growth_pct": 18.1, "tone": "good"},
            ],
            "wallet_benchmark_pct": 12.0, "portfolio_benchmark_pct": 25.0,
            "benchmark_label": "Carrier average / market mix", "basis": "Q2 2026",
        },
        "top_carriers": {
            "carriers": [
                {"carrier": "Zurich", "is_subject": True, "rank": 1, "premium": "£19.2m",
                 "premium_value": 19_200_000, "share_of_wallet_pct": 11.9,
                 "movement": "-14.3% QoQ", "movement_pct": -14.3},
                {"carrier": "Peer 1", "rank": 2, "premium": "£15.0m", "premium_value": 15_000_000,
                 "share_of_wallet_pct": 9.3, "movement": "+4.1% QoQ", "movement_pct": 4.1},
            ],
            "field_size": 12, "scope": "Canada / Property", "basis": "Q2 2026 vs Q1 2026",
        },
    }


def walk(node):
    if isinstance(node, Component):
        yield node
        for child in node._traverse():
            if isinstance(child, Component):
                yield child
    elif isinstance(node, (list, tuple)):
        for item in node:
            yield from walk(item)


def text_of(node) -> str:
    """Every string anywhere in a rendered tree, for content assertions."""
    parts = []
    for component in walk(node):
        for value in vars(component).values():
            if isinstance(value, str):
                parts.append(value)
    return " ".join(parts)


# ── every explainable kind is fully wired ────────────────────────────────────


@pytest.mark.parametrize("kind", EXPLAINABLE_KINDS)
def test_each_kind_has_a_renderer_an_editor_and_a_library_entry(kind):
    """A half-registered widget renders blank or cannot be edited — check all four."""
    assert kind in catalog.GENERATED
    assert kind in catalog.LIBRARY_BY_KIND
    assert kind in editor.KIND_EDITORS
    assert kind in ppt_export._RENDERERS


@pytest.mark.parametrize("kind", EXPLAINABLE_KINDS)
def test_a_library_widget_of_each_kind_renders_and_edits(kind):
    widget = model.make_widget(
        kind, copy.deepcopy(catalog.LIBRARY_BY_KIND[kind]["default_data"]), origin="user"
    )
    assert isinstance(widgets_generated.render_bespoke(kind, widget["data"]), Component)
    assert editor.build_editor_body(widget)


def test_no_explainable_editor_offers_a_severity_or_score_field():
    """The retired widgets let a user type `High` or `70`; these must not."""
    banned = {"severity", "gap_score", "intensity", "premium_strength", "broker_perception"}
    for kind in EXPLAINABLE_KINDS:
        cfg = editor.KIND_EDITORS[kind]
        keys = {f["key"] for f in cfg["scalars"]}
        for list_key in cfg["lists"]:
            keys |= {f["key"] for f in editor.LIST_SPECS[list_key]["fields"]}
        assert not (keys & banned), f"{kind} still edits a score: {keys & banned}"


# ── the document the digest becomes ──────────────────────────────────────────


def test_the_digest_lays_out_as_explainable_pages():
    doc = builder.build_document_from_digest(digest(), n_charts=0)
    kinds = [w["kind"] for page in doc["pages"] for w in page["widgets"]]
    assert set(EXPLAINABLE_KINDS) <= set(kinds)
    # The funnel: widest question first, narrowest last.
    assert [p["title"] for p in doc["pages"]] == [
        "Overview",
        "Pain points",
        "Product deep dive",
        "Industry focus",
    ]
    assert all(p["caption"].endswith("?") for p in doc["pages"]), "each step states its question"


def test_a_watchlist_replaces_the_legacy_severity_bars():
    """Both would be the same risks told two ways; the explainable one wins."""
    doc = builder.build_document_from_digest(digest(), n_charts=0)
    commentary = next(
        w for page in doc["pages"] for w in page["widgets"] if w["kind"] == "commentary"
    )
    assert commentary["data"]["risks"] == []


def test_a_digest_without_a_watchlist_keeps_its_legacy_risks():
    """Saved conversations predate the watchlist and must still render their risks."""
    legacy = {k: v for k, v in digest().items() if k != "watchlist"}
    doc = builder.build_document_from_digest(legacy, n_charts=0)
    commentary = next(
        w for page in doc["pages"] for w in page["widgets"] if w["kind"] == "commentary"
    )
    assert commentary["data"]["risks"][0]["label"] == "Rank decline"


def test_the_funnel_narrows_from_the_book_to_one_industry():
    """A reader should be able to follow overview -> pain -> product -> industry."""
    pages = {p["title"]: [w["kind"] for w in p["widgets"]]
             for p in builder.build_document_from_digest(digest(), 0)["pages"]}
    assert "kpi" in pages["Overview"] and "quarterly" in pages["Overview"]
    assert "watchlist" in pages["Pain points"] and "top_carriers" in pages["Pain points"]
    assert "portfolio_map" in pages["Product deep dive"]
    assert "headroom" in pages["Product deep dive"]
    assert pages["Industry focus"] == ["whitespace"]


def test_quarterly_performance_supersedes_the_annual_timeline():
    payload = digest()
    payload["timeline"] = [{"period": "2024", "title": "Premium +12%"}]
    kinds = [
        w["kind"]
        for page in builder.build_document_from_digest(payload, 0)["pages"]
        for w in page["widgets"]
    ]
    assert "quarterly" in kinds and "timeline" not in kinds


def test_the_document_carries_the_scope_the_answer_was_built_from():
    doc = builder.build_document_from_digest(digest(), n_charts=0)
    assert [chip["value"] for chip in doc["scope"]] == ["Canada", "Property"]


# ── what the board actually shows ────────────────────────────────────────────


def test_the_board_states_its_scope_and_every_comparison_basis():
    rendered = render_document(builder.build_document_from_digest(digest(), 0), [])
    shown = text_of(rendered)
    assert "Canada" in shown and "Property" in shown
    assert "Q2 2026 vs Q1 2026" in shown, "a comparison with no named basis is not defensible"


def test_the_watchlist_states_money_at_risk_instead_of_a_severity():
    """High/Medium/Low was the last unexplained rating on the board."""
    rendered = render_document(builder.build_document_from_digest(digest(), 0), [])
    shown = text_of(rendered)
    assert "£12.4m" in shown, "the premium exposed is what ranks a watch item now"
    assert "Premium fell 14.2%" in shown
    for banned in ("Why High?", "Medium priority", "Unrated"):
        assert banned not in shown, f"the board still shows a severity: {banned}"


def test_watch_items_are_ordered_by_the_premium_they_expose():
    payload = digest()
    payload["watchlist"]["items"] = [
        {"risk": "Small", "premium_exposed_value": 1_000_000, "trigger": "small"},
        {"risk": "Large", "premium_exposed_value": 40_000_000, "trigger": "large"},
    ]
    doc = builder.build_document_from_digest(payload, 0)
    widget = next(w for p in doc["pages"] for w in p["widgets"] if w["kind"] == "watchlist")
    # The renderer takes the order it is given, so ordering has to be applied by
    # the completion layer that the digest and the editor both run.
    completed = derive.complete_watchlist(widget["data"]["watchlist"])
    assert [i["risk"] for i in completed["items"]] == ["Large", "Small"]


def test_the_board_shows_money_where_the_old_widgets_showed_a_score():
    rendered = render_document(builder.build_document_from_digest(digest(), 0), [])
    shown = text_of(rendered)
    assert "£33.8m" in shown  # product-line whitespace
    assert "£18.6m" in shown  # industry whitespace
    assert "No current premium" in shown, "zero premium must read as a fact, not a 100 score"


def test_whitespace_offers_its_product_line_filter():
    rendered = render_document(builder.build_document_from_digest(digest(), 0), [])
    selects = [n for n in walk(rendered) if type(n).__name__ == "Select"]
    assert selects, "Industry Whitespace ranks within a product line, so the filter is required"
    assert any(
        (getattr(n, "className", "") or "") == "bm-x-select" for n in selects
    )
    options = [n for n in walk(rendered) if type(n).__name__ == "Option"]
    assert {str(o.children) for o in options} >= {"Property", "Cyber"}


def test_rows_carry_the_data_attributes_the_client_side_filter_reads():
    rendered = render_document(builder.build_document_from_digest(digest(), 0), [])
    products = {
        getattr(n, "data-product", None)
        for n in walk(rendered)
        if getattr(n, "data-product", None)
    }
    assert products == {"property", "cyber"}


def test_a_widget_the_data_cannot_support_is_dropped_not_explained():
    """A panel that only says "not available here" is a panel to leave out."""
    payload = digest()
    payload["quarterly"] = {"rows": [], "note": "No comparable quarters in this scope."}
    doc = builder.build_document_from_digest(payload, 0)
    kinds = [w["kind"] for page in doc["pages"] for w in page["widgets"]]
    assert "quarterly" not in kinds
    assert "No comparable quarters in this scope." not in text_of(render_document(doc, []))


def test_a_page_whose_every_widget_is_empty_is_dropped_whole():
    payload = digest()
    payload["whitespace"] = {"rows": [], "note": "No industry split in this data."}
    payload.pop("opportunities", None)
    payload.pop("opportunity_map", None)
    doc = builder.build_document_from_digest(payload, 0)
    assert "Industry focus" not in [p["title"] for p in doc["pages"]]


def test_no_widget_kind_is_laid_out_twice():
    doc = builder.build_document_from_digest(digest(), n_charts=2)
    kinds = [
        w["kind"] for page in doc["pages"] for w in page["widgets"] if w["kind"] != "charts"
    ]
    assert len(kinds) == len(set(kinds)), f"a panel is repeated: {kinds}"


# ── editing keeps the derived label true ─────────────────────────────────────


def test_editing_a_premium_re_derives_the_measures_that_follow_from_it():
    """An edited carrier premium changes the whitespace and the share of wallet."""
    doc = builder.build_document_from_digest(digest(), 0)
    widget = next(w for p in doc["pages"] for w in p["widgets"] if w["kind"] == "headroom")
    editor.apply_editor(
        widget,
        {
            "headroom.rows.0.product_line": "Property",
            "headroom.rows.0.carrier_premium_value": "21000000",
            "headroom.rows.0.marsh_premium_value": "42000000",
            "headroom.rows.0.whitespace_premium_value": "",
            "headroom.rows.0.share_of_wallet_pct": "",
        },
        "tester",
    )
    row = widget["data"]["headroom"]["rows"][0]
    assert row["whitespace_premium_value"] == 21_000_000
    assert row["share_of_wallet_pct"] == 50.0


def test_editing_keeps_percentages_intact():
    """A percentage rounded to an int would silently change the reported movement."""
    doc = builder.build_document_from_digest(digest(), 0)
    widget = next(w for p in doc["pages"] for w in p["widgets"] if w["kind"] == "quarterly")
    editor.apply_editor(widget, {"quarterly.rows.0.quarter": "Q1 2026",
                                 "quarterly.rows.0.change_pct": "11.4"}, "tester")
    assert widget["data"]["quarterly"]["rows"][0]["change_pct"] == 11.4


def test_a_blank_measure_stays_unreported_rather_than_becoming_zero():
    doc = builder.build_document_from_digest(digest(), 0)
    widget = next(w for p in doc["pages"] for w in p["widgets"] if w["kind"] == "headroom")
    editor.apply_editor(widget, {"headroom.rows.0.product_line": "Property",
                                 "headroom.rows.0.carrier_premium_value": ""}, "tester")
    assert widget["data"]["headroom"]["rows"][0]["carrier_premium_value"] is None


# ── the exported deck ────────────────────────────────────────────────────────


def test_the_deck_exports_every_explainable_widget():
    doc = builder.build_document_from_digest(digest(), 0)
    blob = ppt_export.export_pptx(doc, [])
    assert blob[:2] == b"PK" and len(blob) > 20_000


def test_the_slide_states_the_same_scope_and_values_as_the_screen():
    from pptx import Presentation
    import io

    doc = builder.build_document_from_digest(digest(), 0)
    prs = Presentation(io.BytesIO(ppt_export.export_pptx(doc, [])))
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    text.extend(cell.text for cell in row.cells)
    deck = " ".join(text)
    assert "Country: Canada | Product: Property" in deck
    assert "£12.4m" in deck, "the exposure on screen must be the exposure on the slide"
    assert "Q2 2026 vs Q1 2026" in deck
    assert "High" not in deck, "the slide must not print a severity the board dropped"


# ── the two share measures ───────────────────────────────────────────────────


def test_every_product_view_shows_share_of_portfolio_beside_share_of_wallet():
    """The two have different denominators, so one without the other misleads."""
    rendered = render_document(builder.build_document_from_digest(digest(), 0), [])
    shown = text_of(rendered)
    assert shown.count("Share of portfolio") >= 3, "product rows must carry the portfolio share"
    assert "Share of wallet" in shown
    assert "46.0%" in shown  # Property's share of the carrier's own book


def test_the_portfolio_map_plots_the_two_shares_and_sizes_by_premium():
    rendered = render_document(builder.build_document_from_digest(digest(), 0), [])
    bubbles = [
        n for n in walk(rendered) if (getattr(n, "className", "") or "") == "bm-x-bubble"
    ]
    assert len(bubbles) == 2
    # Area-proportional: the £8.2m line draws bigger than the £0.4m one.
    sizes = [float((b.style or {})["width"].rstrip("px")) for b in bubbles]
    assert sizes[0] > sizes[1]
    shown = text_of(rendered)
    assert "Share of wallet — how much of Marsh's placement is won" in shown
    assert "Share of portfolio — how the book is split" in shown
    assert "Bubble size = premium in the line" in shown


def test_a_bubble_reads_as_a_business_position_not_a_coordinate():
    rendered = render_document(builder.build_document_from_digest(digest(), 0), [])
    titles = " ".join(getattr(n, "title", "") or "" for n in walk(rendered))
    assert "Core strength" in titles, "above both benchmarks is the carrier's core"
    assert "Low presence" in titles


def test_whitespace_says_whether_the_premium_is_unplaced_or_written_by_peers():
    rendered = render_document(builder.build_document_from_digest(digest(), 0), [])
    shown = text_of(rendered)
    assert "Peers hold" in shown and "34.0%" in shown
    assert "Largest unwritten premium" in shown, "the focus line tells the reader where to go"


# ── top carriers ─────────────────────────────────────────────────────────────


def test_top_carriers_ranks_with_its_field_size_and_marks_the_subject():
    rendered = render_document(builder.build_document_from_digest(digest(), 0), [])
    shown = text_of(rendered)
    assert "#1 of 12" in shown, "a rank without its field size is meaningless"
    assert "This carrier" in shown


def test_top_carriers_keeps_peers_anonymous():
    """Peer identities are redacted at the evidence boundary; nothing re-names them."""
    rendered = render_document(builder.build_document_from_digest(digest(), 0), [])
    assert "Peer 1" in text_of(rendered)


def test_the_slide_notes_carry_the_trigger_behind_a_watch_item():
    from pptx import Presentation
    import io

    doc = builder.build_document_from_digest(digest(), 0)
    prs = Presentation(io.BytesIO(ppt_export.export_pptx(doc, [])))
    notes = " ".join(
        s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide
    )
    assert "Premium fell 14.2%" in notes
    assert "no severity" in notes.lower()


# ── the numbers the model left out ───────────────────────────────────────────


def test_a_product_row_gets_the_shares_its_own_premiums_imply():
    """The model reports what it finds; the arithmetic is not left to it."""
    completed = derive.complete_headroom(
        {
            "rows": [
                {"product_line": "Property", "carrier_premium_value": 8_200_000,
                 "marsh_premium_value": 42_000_000},
                {"product_line": "Cyber", "carrier_premium_value": 1_800_000,
                 "marsh_premium_value": 6_000_000},
            ]
        }
    )
    prop = next(r for r in completed["rows"] if r["product_line"] == "Property")
    assert prop["share_of_wallet_pct"] == 19.5      # 8.2 / 42.0
    assert prop["share_of_portfolio_pct"] == 82.0   # 8.2 / (8.2 + 1.8)
    assert prop["whitespace_premium_value"] == 33_800_000
    assert prop["carrier_premium"] == "£8.2m", "a number must arrive with its display text"


def test_a_row_with_no_premium_at_all_is_dropped():
    completed = derive.complete_headroom(
        {"rows": [{"product_line": "Property", "carrier_premium_value": 8_200_000},
                  {"product_line": "Marine"}]}
    )
    assert [r["product_line"] for r in completed["rows"]] == ["Property"]


def test_the_shares_are_never_derived_from_each_other():
    """Wallet and portfolio have different denominators — a governed distinction."""
    completed = derive.complete_headroom(
        {"rows": [{"product_line": "Property", "carrier_premium_value": 5_000_000}]}
    )
    row = completed["rows"][0]
    assert row["share_of_wallet_pct"] is None, "no Marsh premium means no share of wallet"
    assert row["share_of_portfolio_pct"] == 100.0


def test_a_reported_figure_is_never_overwritten():
    completed = derive.complete_headroom(
        {"rows": [{"product_line": "Property", "carrier_premium_value": 8_200_000,
                   "marsh_premium_value": 42_000_000, "share_of_wallet_pct": 21.0}]}
    )
    assert completed["rows"][0]["share_of_wallet_pct"] == 21.0


def test_the_board_prints_the_share_of_wallet_the_model_omitted():
    payload = digest()
    for row in payload["headroom"]["rows"]:
        row.pop("share_of_wallet_pct", None)
        row.pop("whitespace_premium", None)
        row.pop("whitespace_premium_value", None)
    payload["headroom"] = derive.complete_headroom(payload["headroom"])
    shown = text_of(render_document(builder.build_document_from_digest(payload, 0), []))
    assert "19.5%" in shown
    assert "£33.8m" in shown


# ── the product page says each thing once ───────────────────────────────────


def test_the_product_page_lists_its_lines_once():
    """The map's legend used to repeat Product Line Headroom under the plot."""
    rendered = render_document(builder.build_document_from_digest(digest(), 0), [])
    titles = [
        str(n.children)
        for n in walk(rendered)
        if (getattr(n, "className", "") or "") == "bm-x-row-title"
    ]
    assert titles.count("Property") == 1, f"a product line is listed twice: {titles}"


def test_the_map_keeps_the_plot_and_headroom_keeps_the_numbers():
    rendered = render_document(builder.build_document_from_digest(digest(), 0), [])
    shown = text_of(rendered)
    assert "Product portfolio map" in shown and "Product line headroom" in shown
    # The whitespace money is stated by headroom, not by the map.
    assert "£33.8m" in shown


# ── the bubble plot ─────────────────────────────────────────────────────────


def _bubbles(n: int) -> list:
    return [
        {
            "product_line": f"Line {i}",
            "share_of_wallet_pct": 10.0 + i,
            "share_of_portfolio_pct": 5.0 + i,
            "premium": f"£{10 - i}.0m",
            "premium_value": (10 - i) * 1_000_000,
        }
        for i in range(n)
    ]


def test_the_map_plots_only_the_seven_biggest_lines():
    payload = digest()
    payload["portfolio_map"] = {"bubbles": _bubbles(11), "benchmark_label": "Peer average"}
    rendered = render_document(builder.build_document_from_digest(payload, 0), [])
    plotted = [
        n for n in walk(rendered)
        if (getattr(n, "className", "") or "").startswith("bm-x-bubble-point")
    ]
    assert len(plotted) == 7


def test_the_lines_it_did_not_plot_are_named_not_hidden():
    payload = digest()
    payload["portfolio_map"] = {"bubbles": _bubbles(11), "benchmark_label": "Peer average"}
    shown = text_of(render_document(builder.build_document_from_digest(payload, 0), []))
    assert "4 smaller lines not plotted" in shown
    assert "Line 10" in shown


def test_a_small_book_is_plotted_whole():
    payload = digest()
    payload["portfolio_map"] = {"bubbles": _bubbles(4), "benchmark_label": "Peer average"}
    rendered = render_document(builder.build_document_from_digest(payload, 0), [])
    plotted = [
        n for n in walk(rendered)
        if (getattr(n, "className", "") or "").startswith("bm-x-bubble-point")
    ]
    assert len(plotted) == 4
    assert "not plotted" not in text_of(rendered)


# ── whitespace colour ───────────────────────────────────────────────────────


def test_a_growing_market_the_carrier_is_growing_into_is_the_prime_gap():
    assert opportunity.classify(
        marsh_change_pct=8.2, carrier_change_pct=11.4, share_of_wallet_pct=0.0
    ) == "prime"


def test_a_growing_market_the_carrier_is_losing_reads_differently():
    assert opportunity.classify(
        marsh_change_pct=8.2, carrier_change_pct=-4.0, share_of_wallet_pct=5.9
    ) == "losing"
    assert opportunity.classify(
        marsh_change_pct=8.2, carrier_change_pct=None, share_of_wallet_pct=5.9
    ) == "open"


def test_a_flat_market_or_a_held_wallet_is_not_a_gap():
    assert opportunity.classify(
        marsh_change_pct=0.4, carrier_change_pct=9.0, share_of_wallet_pct=1.0
    ) == "flat"
    assert opportunity.classify(
        marsh_change_pct=8.0, carrier_change_pct=9.0, share_of_wallet_pct=50.0
    ) == "held"


def test_a_row_with_no_market_movement_is_left_uncoloured():
    """A colour nobody can check is decoration; no movement, no band."""
    assert opportunity.classify(
        marsh_change_pct=None, carrier_change_pct=9.0, share_of_wallet_pct=1.0
    ) == ""


def test_the_band_travels_from_the_facts_to_the_row():
    completed = derive.complete_whitespace(
        {
            "rows": [
                {"industry": "Manufacturing", "marsh_premium_value": 18_600_000,
                 "carrier_premium_value": 0.0, "marsh_change_pct": 8.2,
                 "carrier_change_pct": 11.4},
            ]
        }
    )
    row = completed["rows"][0]
    assert row["opportunity"] == "prime"
    # A carrier premium of ZERO is the point of a whitespace row, not a blank.
    assert row["share_of_wallet_pct"] == 0.0
    assert row["whitespace_premium_value"] == 18_600_000


def test_the_board_colours_the_row_and_says_what_the_colour_means():
    payload = digest()
    payload["whitespace"] = derive.complete_whitespace(
        {
            "rows": [
                {"industry": "Manufacturing", "product_line": "Property",
                 "marsh_premium_value": 18_600_000, "carrier_premium_value": 0.0,
                 "marsh_change": "+8.2% QoQ", "marsh_change_pct": 8.2,
                 "carrier_change": "+11.4% QoQ", "carrier_change_pct": 11.4},
            ],
            "product_line": "Property",
            "product_lines": ["Property"],
        }
    )
    rendered = render_document(builder.build_document_from_digest(payload, 0), [])
    bands = [
        n for n in walk(rendered)
        if (getattr(n, "className", "") or "").startswith("bm-x-band")
    ]
    assert bands and "good" in bands[0].className
    shown = text_of(rendered)
    assert "Prime gap" in shown
    assert "What the colours mean" in shown
    # The two numbers the colour is read from are printed beside it.
    assert "+8.2% QoQ" in shown and "+11.4% QoQ" in shown


def test_the_slide_carries_the_same_band_and_the_same_rule():
    from pptx import Presentation
    import io

    payload = digest()
    payload["whitespace"] = derive.complete_whitespace(
        {
            "rows": [
                {"industry": "Manufacturing", "product_line": "Property",
                 "marsh_premium_value": 18_600_000, "carrier_premium_value": 0.0,
                 "marsh_change": "+8.2% QoQ", "marsh_change_pct": 8.2,
                 "carrier_change_pct": 11.4},
            ],
            "product_line": "Property",
        }
    )
    doc = builder.build_document_from_digest(payload, 0)
    prs = Presentation(io.BytesIO(ppt_export.export_pptx(doc, [])))
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text.append(shape.text_frame.text)
    deck = " ".join(text)
    assert "Prime gap" in deck
    assert "Colour reads the two growth rates" in deck

def test_editing_a_movement_re_colours_the_row():
    """The band is derived, so an edited growth rate must re-band, not go stale."""
    payload = digest()
    payload["whitespace"] = derive.complete_whitespace(
        {
            "rows": [
                {"industry": "Manufacturing", "product_line": "Property",
                 "marsh_premium_value": 18_600_000, "carrier_premium_value": 0.0,
                 "marsh_change_pct": 8.2, "carrier_change_pct": 11.4},
            ],
            "product_line": "Property",
        }
    )
    doc = builder.build_document_from_digest(payload, 0)
    widget = next(w for p in doc["pages"] for w in p["widgets"] if w["kind"] == "whitespace")
    assert widget["data"]["whitespace"]["rows"][0]["opportunity"] == "prime"

    editor.apply_editor(
        widget,
        {
            "whitespace.rows.0.industry": "Manufacturing",
            "whitespace.rows.0.marsh_change_pct": "8.2",
            "whitespace.rows.0.carrier_change_pct": "-4.0",
        },
        "tester",
    )
    assert widget["data"]["whitespace"]["rows"][0]["opportunity"] == "losing"


def test_the_editor_collects_both_movements_and_never_the_band():
    cfg = editor.KIND_EDITORS["whitespace"]
    keys = {f["key"] for lk in cfg["lists"] for f in editor.LIST_SPECS[lk]["fields"]}
    assert {"marsh_change_pct", "carrier_change_pct"} <= keys
    assert "opportunity" not in keys, "a derived band is not an author's to type"

def test_the_map_always_has_both_benchmark_axes():
    """Without them the plot is dots with nothing to read them against."""
    completed = derive.complete_portfolio_map(
        {
            "bubbles": [
                {"product_line": "Property", "carrier_premium_value": 8_200_000,
                 "marsh_premium_value": 42_000_000},
                {"product_line": "Cyber", "carrier_premium_value": 1_800_000,
                 "marsh_premium_value": 6_000_000},
            ]
        }
    )
    # Premium-weighted, not a mean of percentages: 10.0m of 48.0m.
    assert completed["wallet_benchmark_pct"] == 20.8
    # An even split of the carrier's own book across its lines.
    assert completed["portfolio_benchmark_pct"] == 50.0


def test_a_reported_benchmark_is_never_overwritten():
    completed = derive.complete_portfolio_map(
        {
            "bubbles": [
                {"product_line": "Property", "carrier_premium_value": 8_200_000,
                 "marsh_premium_value": 42_000_000},
            ],
            "wallet_benchmark_pct": 17.0,
        }
    )
    assert completed["wallet_benchmark_pct"] == 17.0


def test_both_axes_are_drawn_and_named_on_the_board():
    payload = digest()
    payload["portfolio_map"] = derive.complete_portfolio_map(
        {
            "bubbles": [
                {"product_line": "Property", "carrier_premium_value": 8_200_000,
                 "marsh_premium_value": 42_000_000},
                {"product_line": "Cyber", "carrier_premium_value": 1_800_000,
                 "marsh_premium_value": 6_000_000},
            ]
        }
    )
    rendered = render_document(builder.build_document_from_digest(payload, 0), [])
    axes = {
        (getattr(n, "className", "") or "")
        for n in walk(rendered)
        if (getattr(n, "className", "") or "").startswith("bm-x-benchline")
    }
    assert axes == {"bm-x-benchline v", "bm-x-benchline h"}
    shown = text_of(rendered)
    assert "Avg share of wallet 20.8%" in shown
    assert "Avg share of portfolio 50.0%" in shown


def test_the_slide_names_the_same_two_axes():
    from pptx import Presentation
    import io

    payload = digest()
    payload["portfolio_map"] = derive.complete_portfolio_map(
        {
            "bubbles": [
                {"product_line": "Property", "carrier_premium_value": 8_200_000,
                 "marsh_premium_value": 42_000_000},
                {"product_line": "Cyber", "carrier_premium_value": 1_800_000,
                 "marsh_premium_value": 6_000_000},
            ]
        }
    )
    doc = builder.build_document_from_digest(payload, 0)
    prs = Presentation(io.BytesIO(ppt_export.export_pptx(doc, [])))
    deck = " ".join(
        s.text_frame.text for slide in prs.slides for s in slide.shapes if s.has_text_frame
    )
    assert "Avg share of wallet 20.8%" in deck
    assert "Avg share of portfolio 50.0%" in deck

