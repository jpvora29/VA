"""Retyped words keep the deck's formatting — on the canvas and in the exported .pptx.

The regression: editing any text box on the Canvas drew the new words in one plain
style (no bullets, the wrong size and colour), and the export wrote line *i* into
paragraph *i* even though the edit field skips blank spacer paragraphs — so every line
after a spacer landed in the wrong paragraph and took its formatting.
"""
from __future__ import annotations

from copy import deepcopy

import pytest
from dash.development.base_component import Component
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from studio.page import template_preview as TP
from studio.template_fill import text_edits as TE
from studio.template_fill.analyze import analyze
from studio.template_fill.fill import apply_text_overrides


# ── a small deck: heading, spacer, two bullets (one with a bold lead-in) ──────


def _bullet(paragraph, char="•", margin=Emu(228600), indent=Emu(-228600)) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    ppr.set("marL", str(int(margin)))
    ppr.set("indent", str(int(indent)))
    bu = etree.SubElement(ppr, qn("a:buChar"))
    bu.set("char", char)


def _run(paragraph, text, *, size=12, bold=False, rgb="0B1F44"):
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(rgb)
    return run


@pytest.fixture()
def deck(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(457200), Emu(457200), Emu(6000000), Emu(3000000))
    frame = box.text_frame
    _run(frame.paragraphs[0], "Highlights", size=20, bold=True, rgb="C00000")
    frame.add_paragraph()                                  # the template's spacer
    first = frame.add_paragraph()
    _bullet(first)
    _run(first, "Growth: ", bold=True)
    _run(first, "premium rose 12%.")
    second = frame.add_paragraph()
    _bullet(second)
    _run(second, "Share held at 9%.")
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path, int(box.shape_id)


def _paragraphs(path, shape_id):
    shape = next(s for s in Presentation(path).slides[0].shapes if int(s.shape_id) == shape_id)
    return list(shape.text_frame.paragraphs)


def _has_bullet(paragraph) -> bool:
    ppr = paragraph._p.find(qn("a:pPr"))
    return ppr is not None and ppr.find(qn("a:buChar")) is not None


# ── the line → paragraph mapping ─────────────────────────────────────────────


def test_lines_go_back_to_the_worded_paragraphs_they_came_from():
    assigned, extra = TE.assign_lines(["Title", "", "One", "Two"], ["T", "1", "2", "3"])

    assert assigned == [(0, ["T"]), (2, ["1"]), (3, ["2"])]
    assert extra == ["3"]


def test_a_deleted_line_leaves_its_paragraph_empty_handed():
    assigned, extra = TE.assign_lines(["Title", "", "One", "Two"], ["T", "1"])

    assert assigned == [(0, ["T"]), (2, ["1"]), (3, [])]
    assert extra == []


def test_a_new_line_looks_like_the_last_paragraph():
    owners = TE.line_paragraphs(["Title", "", "One"], ["T", "1", "new"])

    assert owners == [(0, True), (2, True), (2, True)]


# ── the export ───────────────────────────────────────────────────────────────


def test_retyping_the_heading_leaves_the_bullets_and_spacer_alone(deck, tmp_path):
    path, shape_id = deck
    out = tmp_path / "out.pptx"
    apply_text_overrides(str(path), {f"0:{shape_id}": [
        "Key highlights", "Growth: premium rose 12%.", "Share held at 9%."]}, str(out))

    paras = _paragraphs(out, shape_id)
    assert [p.text for p in paras] == ["Key highlights", "", "Growth: premium rose 12%.",
                                       "Share held at 9%."]
    assert paras[0].runs[0].font.size == Pt(20) and paras[0].runs[0].font.bold
    assert _has_bullet(paras[2]) and _has_bullet(paras[3])
    # The untouched bullet keeps its bold lead-in and plain body — two runs, as authored.
    assert [(r.text, bool(r.font.bold)) for r in paras[2].runs] == [
        ("Growth: ", True), ("premium rose 12%.", False)]


def test_a_retyped_bullet_is_written_into_the_bullet_not_the_spacer(deck, tmp_path):
    path, shape_id = deck
    out = tmp_path / "out.pptx"
    apply_text_overrides(str(path), {f"0:{shape_id}": [
        "Highlights", "Growth slowed.", "Share held at 9%."]}, str(out))

    paras = _paragraphs(out, shape_id)
    assert [p.text for p in paras] == ["Highlights", "", "Growth slowed.", "Share held at 9%."]
    assert _has_bullet(paras[2])
    assert paras[2].runs[0].font.size == Pt(12)


def test_an_added_line_is_a_bullet_like_the_last_one(deck, tmp_path):
    path, shape_id = deck
    out = tmp_path / "out.pptx"
    apply_text_overrides(str(path), {f"0:{shape_id}": [
        "Highlights", "Growth: premium rose 12%.", "Share held at 9%.", "Cyber doubled."]},
        str(out))

    paras = _paragraphs(out, shape_id)
    assert paras[-1].text == "Cyber doubled."
    assert _has_bullet(paras[-1]) and paras[-1].runs[0].font.size == Pt(12)


# ── what the canvas draws ────────────────────────────────────────────────────


def test_analysis_resolves_each_paragraph_s_look(deck):
    path, shape_id = deck
    shape = next(s for s in analyze(str(path)).slides[0].shapes if s.shape_id == shape_id)
    heading, _spacer, first, _second = shape.para_styles

    assert (heading.size_pt, heading.color, heading.bold, heading.bullet) == (20.0, "C00000", True, "")
    assert (first.size_pt, first.bullet, first.margin_left_pt, first.indent_pt) == (12.0, "•", 18.0, -18.0)


def test_a_theme_colour_with_a_luminance_tweak_resolves_to_hex():
    from studio.template_fill.para_style import ThemeContext, resolve_color

    fill = etree.fromstring(
        '<a:solidFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:schemeClr val="tx1"><a:lumMod val="50000"/><a:lumOff val="50000"/></a:schemeClr>'
        '</a:solidFill>')

    assert resolve_color(fill, ThemeContext(colors={"tx1": "000000"})) == "808080"


def _walk(node):
    if isinstance(node, (list, tuple)):
        for child in node:
            yield from _walk(child)
        return
    if node is None:
        return
    yield node
    if isinstance(node, Component):
        yield from _walk(getattr(node, "children", None))


def test_the_canvas_draws_retyped_lines_in_their_paragraph_s_style(deck, monkeypatch):
    path, shape_id = deck
    monkeypatch.setattr(TP, "public_url_exists", lambda url: True)
    tdoc = {
        "template_path": str(path), "values": {}, "manifest": [], "hidden": [], "order": [0],
        "assembled": True, "background_urls": ["/assets/render-0.png"],
        "text_edits": {f"0:{shape_id}": ["Key highlights", "Growth slowed.", "Share held."]},
    }

    body = TP.template_preview_body(tdoc, {"idx": 0})
    reflection = next(n for n in _walk(body)
                      if "qs-tf-reflect" in str(getattr(n, "className", "")).split())
    lines = [n for n in _walk(reflection)
             if "qs-tf-reflect-line" in str(getattr(n, "className", ""))]

    assert len(lines) == 3
    heading, bullet = lines[0].style, lines[1].style
    assert heading["fontWeight"] == "700" and heading["color"] == "#C00000"
    assert float(heading["fontSize"][:-2]) > float(bullet["fontSize"][:-2])
    marks = [n.children for n in _walk(lines[1]) if getattr(n, "className", "") == "qs-tf-bullet"]
    assert marks == ["•"], "a bullet paragraph keeps its bullet on the canvas"
    assert float(bullet["paddingLeft"][:-2]) > 0, "and its indent"


# ── a KPI tile: a navy table cell on a pale slide ────────────────────────────


@pytest.fixture()
def kpi_deck(tmp_path):
    """One KPI table: a label row and a bold 18pt value row, each cell filled navy, with
    a 0.2" right margin on the value — the Overall page's tiles."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_table(2, 2, Emu(457200), Emu(457200), Emu(2743200), Emu(731520))
    for r, row in enumerate(shape.table.rows):
        for c, cell in enumerate(row.cells):
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string("0B1F44")
            paragraph = cell.text_frame.paragraphs[0]
            _run(paragraph, ("GWP 2025", "SoW 2025")[c] if r == 0 else ("208M", "9.1%")[c],
                 size=9 if r == 0 else 18, bold=True, rgb="CEECFF")
    value = shape.table.cell(1, 0)
    value.margin_right = Emu(182880)
    path = tmp_path / "kpi.pptx"
    prs.save(path)
    return path, int(shape.shape_id)


def test_a_retyped_kpi_draws_on_its_own_tile_not_the_slide(kpi_deck, monkeypatch):
    """Clicking a KPI and retyping it used to paint the whole cell in the SLIDE's paper,
    so light tile text vanished on a pale box over the tile's icon, at the cell's edge."""
    path, shape_id = kpi_deck
    monkeypatch.setattr(TP, "public_url_exists", lambda url: True)
    tdoc = {
        "template_path": str(path), "values": {}, "manifest": [], "hidden": [], "order": [0],
        "assembled": True, "background_urls": ["/assets/render-0.png"],
        "text_edits": {f"0:{shape_id}:1:0": ["209M"]},
    }

    body = TP.template_preview_body(tdoc, {"idx": 0})
    cell = next(n for n in _walk(body) if getattr(n, "className", "").startswith("qs-tf-cell")
                and n.__dict__.get("data-at") == f"0:{shape_id}:1:0")
    assert cell.style["--qs-tf-paper"].upper() == "#0B1F44", "the tile's own fill is the paper"

    reflection = next(n for n in _walk(cell) if str(getattr(n, "className", "")) ==
                      "qs-tf-reflect is-styled")
    assert reflection.style["padding"].split()[1] != reflection.style["padding"].split()[3], \
        "the cell's own margins place the words (0.2in right, default left)"
    line = next(n for n in _walk(reflection) if "qs-tf-reflect-line" in str(n.className))
    assert line.style["fontWeight"] == "700" and line.style["color"] == "#CEECFF"
    paint = next(n for n in _walk(line) if getattr(n, "className", "") == "qs-tf-paint")
    ghost, words = paint.children
    assert (ghost.children, words.children) == ("208M", "209M"), \
        "the paper is sized to the old words and the new, not the whole cell"
