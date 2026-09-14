"""Themed lead-ins: ``Manufacturing growth: the finding``.

A reader scanning a column sees what each bullet is ABOUT before reading it, which is why
the product/country panels that already do this read easily. The shape is offered to the
writer rather than imposed — a column of invented labels ("Performance:", "Overall:") is
harder to read than none — and policed, because the failure it must be kept away from is
the one a flat ban on colons used to prevent: "Momentum: Cyber +97%", a label in front of
a chart caption rather than a sentence.

Pure: no LLM, no database.
"""
from __future__ import annotations

import pytest

from studio.template_fill import commentary as CM
from studio.template_fill import commentary_metrics as M


GOOD = "Manufacturing growth: Marsh-placed premium there rose 18% while the carrier's share fell 2.1 points."
PLAIN = "The carrier grew 12.0% while the Marsh book grew 6.0%."


# ── what counts as a lead-in ────────────────────────────────────────────────


def test_a_named_thing_followed_by_a_whole_sentence_is_a_lead_in():
    assert M.lead_in_issue(GOOD) is None
    assert M.lead_in_label(GOOD) == "Manufacturing growth"


def test_a_bullet_without_a_colon_is_not_a_lead_in_and_is_left_alone():
    assert M.lead_in_issue(PLAIN) is None
    assert M.lead_in_label(PLAIN) == ""


def test_a_colon_deep_inside_a_sentence_is_ordinary_punctuation():
    line = ("This is a long sentence that happens to contain a colon: and it keeps going "
            "well past anything a label could be.")
    assert M.lead_in_issue(line) is None
    assert M.lead_in_label(line) == ""


# ── and what the shape must never decay into ────────────────────────────────


def test_a_label_in_front_of_a_caption_is_refused():
    """The exact pattern the old flat ban on colons existed to stop."""
    issue = M.lead_in_issue("Momentum: Cyber +97%")
    assert issue and "complete sentence" in issue


def test_a_label_that_repeats_the_column_heading_is_refused():
    """A bullet labelled with the heading it already sits under says where it is, not what."""
    for heading in ("Challenges", "Opportunities", "Key messages", "Performance"):
        issue = M.lead_in_issue(f"{heading}: The carrier lost share across three industries.")
        assert issue and "heading" in issue, heading


def test_the_sentence_after_the_label_must_stand_on_its_own():
    issue = M.lead_in_issue("Marine: fell sharply.")
    assert issue is not None


def test_a_malformed_lead_in_does_not_reach_a_slide():
    """The gate, not only the prompt."""
    assert CM._accept(["Momentum: Cyber +97%"], wanted=1, node="t") is None
    assert CM._accept([GOOD], wanted=1, node="t") == GOOD


# ── the writer is offered it, not ordered to use it ─────────────────────────


def test_the_writer_is_offered_the_shape_with_its_three_rules():
    voice = CM.deck_voice("balanced", "Zurich")
    assert "short lead-in label" in voice
    assert "at most four words" in voice
    assert "COMPLETE SENTENCE" in voice
    assert "do not label every bullet" in voice, "a formula produces invented labels"


def test_the_judge_is_told_the_shape_is_allowed():
    """Writer and judge must move together, or the judge deletes what the writer offers."""
    from studio.template_fill import commentary_verify as V

    assert "lead-in label" in V._JUDGE_SYSTEM
    assert "KEEP that shape" in V._JUDGE_SYSTEM


# ── the slide: one weight, whatever the template placeholder carried ────────


def _one_paragraph_frame():
    from pptx import Presentation
    from pptx.util import Inches

    slide = Presentation().slides.add_slide(Presentation().slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    frame = box.text_frame
    frame.text = "placeholder"
    return frame


def test_commentary_never_inherits_a_bold_placeholder():
    """A panel whose example text was typed bold turned every real line bold.

    ``_set_paragraph_text`` reuses the authored run and keeps its ``rPr`` deliberately —
    right for colour and think-cell fields, wrong for weight, because one panel then reads
    bold while the panel beside it reads normal. Commentary is one weight.
    """
    from studio.template_fill import fill

    frame = _one_paragraph_frame()
    frame.paragraphs[0].runs[0].font.bold = True       # the author's bold placeholder
    fill._write_bullets(frame, GOOD)
    assert frame.paragraphs[0].text == GOOD
    assert not any(r.font.bold for r in frame.paragraphs[0].runs),         "a bold placeholder must not make the commentary bold"


def test_a_lead_in_is_plain_text_not_emphasis():
    """The label earns its place by being FIRST, not by being heavy.

    Emphasis was tried and removed: it read as bold noise down a narrow panel, and the
    label is already doing its job by opening the bullet.
    """
    from studio.template_fill import fill

    frame = _one_paragraph_frame()
    fill._write_bullets(frame, GOOD)
    assert frame.paragraphs[0].text == GOOD
    assert not any(r.font.bold for r in frame.paragraphs[0].runs)
