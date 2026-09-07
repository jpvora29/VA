"""Shared PowerPoint drawing kit for the Boardroom export.

Geometry, brand colours and the handful of shape helpers that every widget
renderer draws with. It lives apart from :mod:`ui.boardroom.ppt_export` so the
widget renderers can be split across modules without importing each other.

Nothing here knows what a widget is: these are text boxes, rounded rectangles,
rules, dots and table cells, in inches.
"""
from __future__ import annotations

import math
import re
from typing import Optional

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from ui.boardroom import model


# -- Geometry: 16:9 slide, CSS px <-> inches at 96 px/in ----------------------
SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
MARGIN_IN = 0.45
HEADER_IN = 0.72
PX_PER_IN = 96.0
GUTTER_IN = 14 / PX_PER_IN  # the on-screen grid gap is 14 px

CONTENT_W = SLIDE_W_IN - 2 * MARGIN_IN
CONTENT_H = SLIDE_H_IN - MARGIN_IN - HEADER_IN - 0.25
COL_W = (CONTENT_W - (model.GRID_COLUMNS - 1) * GUTTER_IN) / model.GRID_COLUMNS

FONT = "Segoe UI"
BULLET = "•  "

# Brand / tone colours (mirrors the UI tone classes + boardroom themes).
NAVY = RGBColor(0x0B, 0x13, 0x20)
GRAY = RGBColor(0x5B, 0x65, 0x77)
LIGHT_BORDER = RGBColor(0xD7, 0xDF, 0xEB)
SOFT_BG = RGBColor(0xF4, 0xF7, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

_TONE_HEX = {
    "good": (0x1F, 0x9D, 0x55),
    "warn": (0xB9, 0x81, 0x0A),
    "danger": (0xC5, 0x35, 0x32),
    "neutral": (0x0B, 0x4B, 0xFF),
}


def _tone(value: Optional[str]) -> str:
    v = (value or "neutral").strip().lower()
    return v if v in _TONE_HEX else "neutral"


def _tone_color(value: Optional[str]) -> RGBColor:
    return RGBColor(*_TONE_HEX[_tone(value)])


def _tone_soft(value: Optional[str], alpha: float = 0.12) -> RGBColor:
    """Tone blended towards white - the PPT stand-in for rgba(tone, alpha)."""
    r, g, b = _TONE_HEX[_tone(value)]
    mix = lambda c: int(round(255 + (c - 255) * alpha))  # noqa: E731
    return RGBColor(mix(r), mix(g), mix(b))


def _hex_rgb(value: str, fallback: RGBColor = NAVY) -> RGBColor:
    m = re.fullmatch(r"#?([0-9a-fA-F]{6})", (value or "").strip())
    return RGBColor.from_string(m.group(1)) if m else fallback


# -- Low-level shape helpers ---------------------------------------------------


def _textbox(slide, x, y, w, h):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    return box, tf


def _para(
    tf,
    text: str,
    *,
    size: float = 11,
    bold: bool = False,
    italic: bool = False,
    color: RGBColor = NAVY,
    align=PP_ALIGN.LEFT,
    bullet: bool = False,
    space_after: float = 2,
    first: bool = False,
):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = (BULLET + text) if bullet else text
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    return p


def _rounded(slide, x, y, w, h, *, fill: Optional[RGBColor] = WHITE,
             line: Optional[RGBColor] = LIGHT_BORDER, radius: float = 0.08):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    try:  # corner radius (fraction of the smaller side)
        shp.adjustments[0] = radius
    except Exception:  # noqa: BLE001 - cosmetic only
        pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


def _rect(slide, x, y, w, h, fill: RGBColor, line: Optional[RGBColor] = None):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


def _dot(slide, x, y, d, fill: RGBColor, line: Optional[RGBColor] = None):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.0)
    shp.shadow.inherit = False
    return shp


def _set_cell(cell, text: str, *, size=10, bold=False, color: RGBColor = NAVY,
              fill: Optional[RGBColor] = None, align=PP_ALIGN.LEFT):
    cell.margin_left = Inches(0.06)
    cell.margin_right = Inches(0.06)
    cell.margin_top = cell.margin_bottom = Inches(0.02)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill


def text_lines(text: str, w_in: float, size_pt: float = 11) -> int:
    """How many wrapped lines a string needs at a given width and size."""
    chars_per_line = max(10, int(w_in * 96 / (size_pt * 0.55)))
    lines = 0
    for para in str(text or "").splitlines() or [""]:
        lines += max(1, math.ceil(len(para) / chars_per_line))
    return lines
